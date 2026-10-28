"""The self-check of exports and the 'no data' note."""

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from rfp.docx_blocks import iter_blocks
from rfp.export import ExportVerificationError, export_answers, verify_docx, verify_xlsx
from rfp.export.formatting import NO_DATA_LABEL
from rfp.export.xlsx_export import AnswerHeader
from rfp.models import Answer, Location, Question
from rfp.questions import extract_xlsx_questions
from rfp.web.common import ru

QUESTIONS = {
    "ru": "Как хранятся резервные копии?",
    "kk": "Деректер қалай сақталады?",
    "en": "How are backups stored?",
}


def answer(qid: str, text: str = "Daily.", status="answered") -> Answer:
    return Answer(qid=qid, text=text, status=status, confidence=0.9)


def workbook_with_extras(path: Path) -> Path:
    """Two sheets, a formula, a merged range and a drop-down list: things an export must not disturb."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["ID", "Question", "Answer"])
    for number, text in enumerate(QUESTIONS.values(), 1):
        sheet.append([f"Q{number}", text, None])
    sheet["E1"] = "=1+1"
    sheet.merge_cells("F1:G1")
    validation = DataValidation(type="list", formula1='"Yes,No"')
    validation.add("D2:D4")
    sheet.add_data_validation(validation)
    workbook.create_sheet("Notes")["A1"] = "Keep me"
    workbook.save(path)
    return path


def exported(tmp_path: Path, statuses: dict[str, str] | None = None):
    source = workbook_with_extras(tmp_path / "in.xlsx")
    questions = extract_xlsx_questions(source, "Sheet1", "B", "C", 1, "A")
    statuses = statuses or {}
    answers = {q.qid: answer(q.qid, status=statuses.get(q.qid, "answered")) for q in questions}
    dest = tmp_path / "out.xlsx"
    export_answers(source, questions, answers, dest)
    return source, dest, questions, answers


# --- no data ------------------------------------------------------------------------------------


def test_no_data_is_a_note_in_the_question_language_on_yellow(tmp_path):
    _, dest, _, _ = exported(tmp_path, {"q1": "no_data", "q2": "no_data", "q3": "no_data"})

    sheet = load_workbook(dest)["Sheet1"]
    assert [sheet[f"C{row}"].value for row in (2, 3, 4)] == [
        NO_DATA_LABEL["ru"],
        NO_DATA_LABEL["kk"],
        NO_DATA_LABEL["en"],
    ]
    assert all(sheet[f"C{row}"].fill.fgColor.rgb.endswith("FFF2CC") for row in (2, 3, 4))


def test_answered_cells_are_not_highlighted(tmp_path):
    _, dest, _, _ = exported(tmp_path, {"q1": "no_data"})

    sheet = load_workbook(dest)["Sheet1"]
    assert sheet["C3"].fill.fgColor.rgb != sheet["C2"].fill.fgColor.rgb
    assert sheet["C3"].value == "Daily."


def test_docx_no_data_is_italic_in_the_question_language(tmp_path):
    source = tmp_path / "in.docx"
    doc = Document()
    for text in QUESTIONS.values():
        doc.add_paragraph(text)
    doc.save(source)
    questions = [
        Question(qid=f"q{n}", text=text, location=Location(block=n - 1))
        for n, text in enumerate(QUESTIONS.values(), 1)
    ]
    answers = {q.qid: answer(q.qid, status="no_data") for q in questions}
    dest = tmp_path / "out.docx"

    export_answers(source, questions, answers, dest)

    paragraphs = Document(dest).paragraphs
    assert [p.text for p in paragraphs[1::2]] == [NO_DATA_LABEL[lang] for lang in ("ru", "kk", "en")]
    assert all(p.runs[0].italic for p in paragraphs[1::2])
    assert [p.text for p in paragraphs[0::2]] == list(QUESTIONS.values())


# --- the check passes on a good export ---------------------------------------------------------


def test_good_xlsx_export_passes_and_an_answer_cannot_become_a_formula(tmp_path):
    source = workbook_with_extras(tmp_path / "in.xlsx")
    questions = extract_xlsx_questions(source, "Sheet1", "B", "C", 1, "A")
    answers = {q.qid: answer(q.qid, '=HYPERLINK("http://x")') for q in questions}
    dest = tmp_path / "out.xlsx"

    export_answers(source, questions, answers, dest)

    assert load_workbook(dest)["Sheet1"]["C2"].data_type == "s"


def test_new_answer_column_is_allowed_to_grow_the_sheet(tmp_path):
    source = tmp_path / "in.xlsx"
    workbook = Workbook()
    workbook.active.title = "S"
    workbook.active.append(["Question"])
    workbook.active.append(["How are backups stored?"])
    workbook.save(source)
    questions = extract_xlsx_questions(source, "S", "A", "B", 1)
    dest = tmp_path / "out.xlsx"

    export_answers(source, questions, {"q1": answer("q1")}, dest, [AnswerHeader("S", 1, "B", "Answer")])

    sheet = load_workbook(dest)["S"]
    assert (sheet["B1"].value, sheet["B2"].value) == ("Answer", "Daily.")


# --- the check catches damage ------------------------------------------------------------------


def tamper(dest: Path, change) -> None:
    workbook = load_workbook(dest)
    change(workbook)
    workbook.save(dest)


@pytest.mark.parametrize(
    ("change", "message"),
    [
        (lambda wb: wb["Sheet1"].__setitem__("B2", "Another question?"), "cell B2: content changed"),
        (lambda wb: wb["Sheet1"].__setitem__("C3", "Wrong"), "not at C3"),
        (lambda wb: wb["Sheet1"].__setitem__("C3", None), "not at C3"),
        (lambda wb: wb["Sheet1"].__setitem__("H9", "stray"), "size is"),
        (lambda wb: wb["Sheet1"].__setitem__("E2", "=2+2"), "number of formulas"),
        (lambda wb: wb["Sheet1"].unmerge_cells("F1:G1"), "merged cells changed"),
        (lambda wb: wb["Sheet1"].data_validations.dataValidation.clear(), "drop-down lists"),
        (lambda wb: wb.remove(wb["Notes"]), "sheets differ"),
        (lambda wb: wb["Notes"].__setitem__("A1", "Changed"), "cell A1: content changed"),
    ],
)
def test_damaged_xlsx_is_refused(tmp_path, change, message):
    source, dest, questions, answers = exported(tmp_path)
    tamper(dest, change)

    with pytest.raises(ExportVerificationError, match=message):
        verify_xlsx(source, dest, questions, answers)


def docx_with_table(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph(QUESTIONS["en"])
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = QUESTIONS["ru"]
    doc.save(path)
    return path


def docx_questions() -> list[Question]:
    return [
        Question(qid="q1", text=QUESTIONS["en"], location=Location(block=0)),
        Question(
            qid="q2",
            text=QUESTIONS["ru"],
            location=Location(block=1, table_row=0, table_col=0, answer_table_row=0, answer_table_col=1),
        ),
    ]


def test_docx_answers_go_under_the_question_and_into_the_answer_cell(tmp_path):
    source = docx_with_table(tmp_path / "in.docx")
    questions = docx_questions()
    answers = {q.qid: answer(q.qid) for q in questions}
    dest = tmp_path / "out.docx"

    export_answers(source, questions, answers, dest)  # includes the check

    doc = Document(dest)
    assert [p.text for p in doc.paragraphs] == [QUESTIONS["en"], "Daily."]
    table = doc.tables[0]
    assert (table.cell(0, 0).text, table.cell(0, 1).text) == (QUESTIONS["ru"], "Daily.")


def _damaged_docx(tmp_path: Path, change) -> tuple[Path, Path, list[Question], dict[str, Answer]]:
    source = docx_with_table(tmp_path / "in.docx")
    questions = docx_questions()
    answers = {q.qid: answer(q.qid) for q in questions}
    dest = tmp_path / "out.docx"
    export_answers(source, questions, answers, dest)
    doc = Document(dest)
    change(doc)
    doc.save(dest)
    return source, dest, questions, answers


@pytest.mark.parametrize(
    "change",
    [
        lambda doc: setattr(doc.paragraphs[0].runs[0], "text", "Changed?"),
        lambda doc: doc.paragraphs[1]._p.getparent().remove(doc.paragraphs[1]._p),
        lambda doc: setattr(doc.tables[0].cell(0, 0).paragraphs[0].runs[0], "text", "Changed?"),
        lambda doc: setattr(doc.tables[0].cell(0, 1).paragraphs[0].runs[0], "text", "Other"),
        lambda doc: doc.tables[0].add_row(),
    ],
)
def test_damaged_docx_is_refused(tmp_path, change):
    source, dest, questions, answers = _damaged_docx(tmp_path, change)

    with pytest.raises(ExportVerificationError):
        verify_docx(source, dest, questions, answers)


def test_block_count_of_a_good_docx_matches(tmp_path):
    source = docx_with_table(tmp_path / "in.docx")
    dest = tmp_path / "out.docx"
    questions = docx_questions()

    export_answers(source, questions, {q.qid: answer(q.qid) for q in questions}, dest)

    assert len(list(iter_blocks(Document(dest)))) == len(list(iter_blocks(Document(source)))) + 1


def test_failed_check_message_is_translated():
    message = (
        "Export check failed: sheet 'Sheet1', cell B2: content changed"
        " | sheets differ: expected ['a'], got ['b']"
    )

    text = ru(message)

    assert text.startswith("Самопроверка выгрузки не прошла, файл не отдан.")
    assert "лист «Sheet1», ячейка B2: содержимое изменилось" in text
    assert "набор листов изменился" in text
