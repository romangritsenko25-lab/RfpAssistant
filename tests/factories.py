"""Builders that generate small DOCX / XLSX / PDF files for tests."""

from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

FONT_CANDIDATES = (
    Path("C:/Windows/Fonts/arial.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/dejavu/DejaVuSans.ttf"),
    Path("/Library/Fonts/Arial.ttf"),
)

LONG_PARAGRAPH = (
    "Our evaluation team has long experience in the regional market and values transparent "
    "cooperation with every vendor, so the whole procedure is documented in detail. "
    "Can you share references from three similar projects? "
    "References will be reviewed during the second stage of the selection by the committee."
)
LONG_QUESTION = "Can you share references from three similar projects?"

# Substrings that a correct extraction must find, one per question or requirement in the test RFP.
RFP_EXPECTED = (
    "Describe your information security policy",
    "Do you perform annual penetration tests",
    "Provide the date of the last audit",
    "must support single sign-on",
    "Укажите, как организовано резервное копирование",
    "Есть ли у вас сертификат ISO 27001",
    "Деректерді шифрлау қалай жүзеге асырылады",
    "Жүйе тәулік бойы жұмыс істеуі тиіс",
    "Do you have a disaster recovery plan",
    "Опишите процесс обработки инцидентов",
    "Please describe your onboarding process",
    "shall provide audit logs",
    "shall support role-based access",
    "must be stored in Kazakhstan",
    LONG_QUESTION,
    "Does the platform support SAML",
    "Укажите срок хранения логов",
)

RFP_NOISE = (
    "Acme Corp is a regional retailer",
    "The purpose of this document",
    "Company founded in 2012",
    "Evaluation criteria are described",
    "Support hours: 24/7",
    "Proposals are accepted until",
    "Компания работает на рынке более десяти лет",
)


def find_unicode_font() -> Path | None:
    return next((path for path in FONT_CANDIDATES if path.exists()), None)


def _set_numbering(paragraph: Paragraph, num_id: int, level: int = 0) -> None:
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = level
    num_pr.get_or_add_numId().val = num_id


def build_rfp_docx(path: Path) -> Path:
    """A synthetic RFP: questions in many styles mixed with ordinary prose."""
    doc = Document()
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("Acme Corp is a regional retailer founded in 2012 and operates 40 stores.")
    doc.add_paragraph("The purpose of this document is to select a vendor for the new platform.")
    doc.add_paragraph("1. Company founded in 2012.")
    doc.add_paragraph("Компания работает на рынке более десяти лет и имеет сотни клиентов.")
    doc.add_heading("2. Questions", level=1)
    for text in (
        "1. Describe your information security policy.",
        "1.1 Do you perform annual penetration tests?",
        "1.2 Provide the date of the last audit.",
        "2) The vendor must support single sign-on (SSO).",
        "3. Укажите, как организовано резервное копирование данных.",
        "4. Есть ли у вас сертификат ISO 27001?",
        "5. Деректерді шифрлау қалай жүзеге асырылады?",
        "6. Жүйе тәулік бойы жұмыс істеуі тиіс.",
    ):
        doc.add_paragraph(text)
    doc.add_paragraph("Evaluation criteria are described in appendix A.")
    for text in (
        "Do you have a disaster recovery plan?",
        "Опишите процесс обработки инцидентов.",
        "Please describe your onboarding process.",
    ):
        doc.add_paragraph(text)
    doc.add_heading("3. Technical requirements", level=1)
    for text in (
        "The solution shall provide audit logs.",
        "The solution shall support role-based access.",
        "Data must be stored in Kazakhstan.",
    ):
        _set_numbering(doc.add_paragraph(text), num_id=1)
    doc.add_paragraph(LONG_PARAGRAPH)
    table = doc.add_table(rows=4, cols=3)
    rows = (
        ("No.", "Requirement", "Response"),
        ("1", "Does the platform support SAML?", ""),
        ("2", "Укажите срок хранения логов.", ""),
        ("3", "Support hours: 24/7", ""),
    )
    for table_row, values in zip(table.rows, rows, strict=True):
        for cell, value in zip(table_row.cells, values, strict=True):
            cell.text = value
    doc.add_paragraph("Proposals are accepted until the end of the month.")
    doc.save(path)
    return path


def build_questionnaire_xlsx(path: Path) -> Path:
    """Styled questionnaire: header on row 2, a blank row between questions, a second sheet."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Questionnaire"
    sheet["A1"] = "Security questionnaire"
    for col, title in zip("ABC", ("ID", "Question", "Answer"), strict=True):
        sheet[f"{col}2"] = title
        sheet[f"{col}2"].font = Font(bold=True)
        sheet[f"{col}2"].fill = PatternFill("solid", fgColor="DDDDDD")
    sheet["A3"], sheet["B3"] = "IAM-01", "Do you enforce multi-factor authentication?"
    sheet["A4"], sheet["B4"] = "IAM-02", "Как часто пересматриваются права доступа?"
    sheet["B6"] = "Деректерді сақтау мерзімі қандай?"
    sheet["A6"] = "BCR-01"
    sheet["C3"].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet.column_dimensions["B"].width = 60
    sheet.column_dimensions["C"].width = 45
    workbook.create_sheet("Notes")["A1"] = "Keep me"
    workbook.properties.title = "Security questionnaire"
    workbook.save(path)
    return path


def build_pdf(path: Path, pages: list[list[str]], unicode_font: bool = False) -> Path:
    """One list of text lines per page; ``unicode_font`` embeds a TTF so Cyrillic is extractable."""
    font = "Helvetica"
    if unicode_font:
        font_path = find_unicode_font()
        assert font_path is not None
        font = "TestSans"
        if font not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(font, str(font_path)))
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setTitle("Test PDF")
    for lines in pages:
        pdf.setFont(font, 11)
        y = 800
        for line in lines:
            pdf.drawString(50, y, line)
            y -= 18
        pdf.showPage()
    pdf.save()
    return path
