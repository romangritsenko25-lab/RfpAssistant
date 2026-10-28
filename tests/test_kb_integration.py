"""Parsers against the generated knowledge base and the dev question set (run in CI with ``-m kb``)."""

import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from rfp.chunking import chunk_document
from rfp.documents import parse_document
from rfp.questions import extract_docx_questions, extract_xlsx_questions

DATA = Path(__file__).resolve().parents[1] / "data"
KB = DATA / "kb"
DEV = DATA / "sets" / "dev"

pytestmark = [
    pytest.mark.kb,
    pytest.mark.skipif(not (KB / "manifest.json").exists(), reason="data/kb has not been built"),
]


def manifest_entries() -> list[dict[str, str]]:
    return json.loads((KB / "manifest.json").read_text("utf-8"))["documents"]


def source_headings(doc_id: str) -> set[str]:
    text = (DATA / "kb_src" / f"{doc_id}.md").read_text("utf-8")
    return set(re.findall(r"^## (.+?)\s*$", text, re.MULTILINE))


def gold(file: str) -> list[dict]:
    lines = (DEV / "gold.jsonl").read_text("utf-8").splitlines()
    return [item for item in map(json.loads, filter(None, lines)) if item["file"] == file]


def recall(expected: list[str], found: list[str]) -> float:
    haystack = [text.casefold() for text in found]
    hits = sum(any(item[:30].casefold() in text for text in haystack) for item in expected)
    return hits / len(expected)


@pytest.mark.parametrize("entry", manifest_entries(), ids=lambda e: e["file"])
def test_document_metadata_matches_manifest(entry: dict[str, str]) -> None:
    doc = parse_document(KB / entry["file"])

    assert doc.doc_id == entry["doc_id"]
    assert doc.language == entry["language"]
    assert str(doc.effective_date) == entry["effective_date"]
    assert doc.sha256 == entry["sha256"]
    assert doc.sections
    assert chunk_document(doc)


@pytest.mark.parametrize(
    "entry", [e for e in manifest_entries() if e["format"] != "xlsx"], ids=lambda e: e["file"]
)
def test_sections_cover_source_headings(entry: dict[str, str]) -> None:
    doc = parse_document(KB / entry["file"])

    missing = source_headings(entry["doc_id"]) - {section.heading for section in doc.sections}
    assert not missing


def test_dev_rfp_questions_found() -> None:
    expected = [item["question"] for item in gold("rfp.docx")]
    found = [question.text for question in extract_docx_questions(DEV / "rfp.docx")]

    assert recall(expected, found) >= 0.95
    assert len(found) <= len(expected) + 1


def test_dev_questionnaire_questions_all_found() -> None:
    path = DEV / "questionnaire.xlsx"
    sheet = load_workbook(path, read_only=True).worksheets[0]
    columns = {cell.value: cell.column_letter for cell in next(sheet.iter_rows(max_row=1))}

    questions = extract_xlsx_questions(
        path, sheet.title, columns["Question"], columns["Answer"], number_col=columns["ID"]
    )

    expected = [item["question"] for item in gold("questionnaire.xlsx")]
    assert recall(expected, [q.text for q in questions]) == 1.0
