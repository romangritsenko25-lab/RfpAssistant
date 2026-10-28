from datetime import date, datetime
from pathlib import Path

import pytest
from docx import Document as new_docx
from openpyxl import Workbook
from pypdf import PdfWriter
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from rfp.documents import parse_document
from rfp.documents.language import detect_language, normalize_language
from tests import factories


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("Information security policy of the company", "en"),
        ("Политика информационной безопасности компании", "ru"),
        ("Компанияның ақпараттық қауіпсіздік саясаты және құпиялылық", "kk"),
        ("", "en"),
    ],
)
def test_detect_language_by_alphabet(text: str, expected: str) -> None:
    assert detect_language(text) == expected


def test_normalize_language_keeps_two_letter_code() -> None:
    assert normalize_language("ru-RU") == "ru"
    assert normalize_language("  ") is None


def test_docx_is_split_by_headings_with_metadata(tmp_path: Path) -> None:
    path = tmp_path / "Security Policy.docx"
    doc = new_docx()
    doc.core_properties.title = "Security policy"
    doc.core_properties.language = "en-US"
    doc.core_properties.modified = datetime(2026, 5, 1)
    doc.add_paragraph("Preface text.")
    doc.add_heading("Access control", level=1)
    doc.add_paragraph("Access is granted by role.")
    doc.add_heading("Passwords", level=2)
    doc.add_paragraph("Passwords are rotated yearly.")
    doc.add_heading("Backups", level=1)
    doc.add_paragraph("Backups are daily.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "RPO"
    table.rows[0].cells[1].text = "24h"
    doc.save(path)

    parsed = parse_document(path)

    assert parsed.doc_id == "Security Policy"  # file stem, case kept, when the file carries no id
    assert (parsed.title, parsed.language, parsed.effective_date) == (
        "Security policy",
        "en",
        date(2026, 5, 1),
    )
    assert [s.heading for s in parsed.sections] == [
        "Security policy",
        "Access control",
        "Passwords",
        "Backups",
    ]
    assert [s.order for s in parsed.sections] == [0, 1, 2, 3]
    assert "RPO | 24h" in parsed.sections[-1].text
    assert len(parsed.sha256) == 64


def test_docx_identifier_property_becomes_doc_id(tmp_path: Path) -> None:
    path = tmp_path / "file-name.docx"
    doc = new_docx()
    doc.core_properties.identifier = "SEC-003"
    doc.add_paragraph("Text.")
    doc.save(path)

    assert parse_document(path).doc_id == "SEC-003"


def test_docx_language_falls_back_to_detection(tmp_path: Path) -> None:
    path = tmp_path / "kk.docx"
    doc = new_docx()
    doc.add_heading("Қауіпсіздік", level=1)
    doc.add_paragraph("Деректерді шифрлау және қорғау қағидалары қолданылады.")
    doc.save(path)

    assert parse_document(path).language == "kk"


def test_markdown_front_matter_and_headings(tmp_path: Path) -> None:
    path = tmp_path / "sla.md"
    path.write_text(
        "---\ntitle: SLA\ndoc_id: SLA-001\nlanguage: ru\neffective_date: 2026-02-10\n---\n"
        "# Доступность\nСервис доступен 99.9%.\n## Исключения\nПлановые работы.\n"
        "```\n# not a heading\n```\n# Поддержка\nКруглосуточно.\n",
        encoding="utf-8",
    )

    parsed = parse_document(path)

    assert (parsed.doc_id, parsed.title, parsed.language) == ("SLA-001", "SLA", "ru")
    assert parsed.effective_date == date(2026, 2, 10)
    assert [s.heading for s in parsed.sections] == ["Доступность", "Исключения", "Поддержка"]
    assert "# not a heading" in parsed.sections[1].text


def test_markdown_without_front_matter_uses_file_name_and_detects_language(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text("# Intro\nPlain English text about delivery.\n", encoding="utf-8")

    parsed = parse_document(path)

    assert (parsed.doc_id, parsed.title, parsed.language, parsed.effective_date) == (
        "notes",
        "notes",
        "en",
        None,
    )


def test_xlsx_rows_become_sections_with_header_pairs(tmp_path: Path) -> None:
    path = tmp_path / "prices.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Services"
    sheet.append(["Service", "SLA", "Note"])
    sheet.append(["Hosting", "99.9%", None])
    sheet.append([])
    sheet.append(["Support", "24/7", "Phone"])
    workbook.save(path)

    parsed = parse_document(path)

    assert [s.heading for s in parsed.sections] == ["Services / row 2", "Services / row 4"]
    assert parsed.sections[0].text == "Service: Hosting\nSLA: 99.9%"
    assert parsed.sections[1].text == "Service: Support\nSLA: 24/7\nNote: Phone"


def test_pdf_without_headings_gives_one_section_per_page(tmp_path: Path) -> None:
    path = factories.build_pdf(
        tmp_path / "plain.pdf",
        [["Data is encrypted at rest.", "Keys are rotated."], ["Backups run nightly."]],
    )

    parsed = parse_document(path)

    assert [s.heading for s in parsed.sections] == ["Page 1", "Page 2"]
    assert "Keys are rotated." in parsed.sections[0].text
    assert parsed.title == "Test PDF"
    assert parsed.effective_date is None  # the PDF creation date is never trusted


def test_pdf_with_headings_is_split_by_heading(tmp_path: Path) -> None:
    path = factories.build_pdf(
        tmp_path / "policy.pdf",
        [
            ["1. Access control", "Access is granted by role.", "1.1 Passwords", "Rotated yearly."],
            ["2. BACKUPS", "Backups run nightly."],
        ],
    )

    parsed = parse_document(path)

    assert [s.heading for s in parsed.sections] == ["1. Access control", "1.1 Passwords", "2. BACKUPS"]


def test_pdf_headings_are_recognised_by_font_size(tmp_path: Path) -> None:
    path = tmp_path / "sized.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    y = 800
    for text, size in (
        ("Policy title", 18),
        ("Access", 14),
        ("Access is granted by role.", 10),
        ("Roles are reviewed yearly.", 10),
        ("Backups", 14),
        ("Backups run nightly.", 10),
    ):
        pdf.setFont("Helvetica", size)
        pdf.drawString(50, y, text)
        y -= 24
    pdf.save()

    parsed = parse_document(path)

    assert [s.heading for s in parsed.sections] == ["Access", "Backups"]
    assert parsed.sections[0].text == "Access is granted by role.\nRoles are reviewed yearly."


def test_pdf_custom_info_keys_give_identity(tmp_path: Path) -> None:
    source = factories.build_pdf(tmp_path / "plain.pdf", [["Data is encrypted."]])
    writer = PdfWriter(clone_from=str(source))
    writer.add_metadata({"/DocId": "SEC-009", "/Language": "kk", "/EffectiveDate": "2026-04-15"})
    path = tmp_path / "tagged.pdf"
    with path.open("wb") as handle:
        writer.write(handle)

    parsed = parse_document(path)

    assert (parsed.doc_id, parsed.language, parsed.effective_date) == ("SEC-009", "kk", date(2026, 4, 15))


@pytest.mark.usefixtures("unicode_font")
def test_pdf_kazakh_text_is_extracted_and_language_detected(tmp_path: Path) -> None:
    path = factories.build_pdf(
        tmp_path / "kk.pdf", [["Деректерді шифрлау қағидалары қолданылады."]], unicode_font=True
    )

    parsed = parse_document(path)

    assert "шифрлау" in parsed.sections[0].text
    assert parsed.language == "kk"


def test_unsupported_extension_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "data.txt"
    path.write_text("x")

    with pytest.raises(ValueError, match="Unsupported"):
        parse_document(path)


def test_same_content_gives_same_hash(tmp_path: Path) -> None:
    first, second = tmp_path / "a.md", tmp_path / "b.md"
    first.write_text("# T\ntext", encoding="utf-8")
    second.write_text("# T\ntext", encoding="utf-8")

    assert parse_document(first).sha256 == parse_document(second).sha256
