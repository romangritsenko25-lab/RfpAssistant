"""Jobs with a model-made layout map: stored params, request log, export with the self-check."""

import pytest
from openpyxl import load_workbook

import rfp.export as export_module
from rfp.export import ExportVerificationError
from rfp.export.formatting import NO_DATA_LABEL
from rfp.services import jobs
from rfp.services.errors import ConflictError
from rfp.web.common import ru
from tests.test_layout_llm import SECURITY_MAP, FakeLayoutLLM, security_workbook

pytestmark = pytest.mark.db


@pytest.fixture
def llm_settings(api_settings, tmp_path):
    prompt = tmp_path / "layout.md"
    prompt.write_text("LAYOUT PROMPT", encoding="utf-8")
    return api_settings.model_copy(
        update={"layout_mode": "llm", "llm_api_key": "test-key", "layout_prompt_path": prompt}
    )


def create(db_conn, path, settings, llm):
    return jobs.create_job(
        db_conn, "tester", path.name, path.read_bytes(), jobs.JobParams(), settings=settings, layout_llm=llm
    )


def test_model_map_is_stored_with_the_job_and_logged(db_conn, tmp_path, llm_settings):
    path = security_workbook(tmp_path / "s.xlsx")

    job = create(db_conn, path, llm_settings, FakeLayoutLLM(SECURITY_MAP))

    detail = jobs.get_job(db_conn, job.id)
    assert detail.total == 3
    assert detail.params.layout.source == "llm"
    assert (detail.params.sheet, detail.params.question_col, detail.params.answer_col) == (
        "Security",
        "B",
        "D",
    )
    row = db_conn.execute("SELECT job_id, qid, model, prompt_tokens, error FROM llm_log").fetchone()
    assert row == (job.id, "layout", "fake", 10, None)


def test_rejected_map_is_logged_and_rules_decide(db_conn, tmp_path, llm_settings, questionnaire_xlsx):
    bad = {"sheets": [{"sheet": "Questionnaire", "question_col": "B", "answer_targets": {"text_col": "B"}}]}

    job = create(db_conn, questionnaire_xlsx, llm_settings, FakeLayoutLLM(bad))

    detail = jobs.get_job(db_conn, job.id)
    assert detail.params.layout.source == "rules" and detail.params.layout.note
    assert db_conn.execute("SELECT error FROM llm_log").fetchone()[0]


def test_rules_mode_never_calls_the_model(db_conn, questionnaire_xlsx, api_settings):
    llm = FakeLayoutLLM(SECURITY_MAP)

    create(db_conn, questionnaire_xlsx, api_settings, llm)

    assert llm.requests == []


def test_export_of_a_multi_sheet_model_layout_passes_the_check(db_conn, tmp_path, llm_settings):
    path = security_workbook(tmp_path / "s.xlsx")
    job = create(db_conn, path, llm_settings, FakeLayoutLLM(SECURITY_MAP))
    db_conn.execute(
        "UPDATE job_items SET status = 'answered', text = 'Yes, daily.', confidence = 0.9 WHERE job_id = %s",
        (job.id,),
    )
    db_conn.execute("UPDATE job_items SET status = 'no_data' WHERE job_id = %s AND qid = 'q3'", (job.id,))
    db_conn.execute("UPDATE jobs SET status = 'done' WHERE id = %s", (job.id,))

    exported = jobs.export_job(db_conn, job.id, "tester")

    out = tmp_path / "out.xlsx"
    out.write_bytes(exported.content)
    sheet = load_workbook(out)["Security"]
    assert sheet["D3"].value == "Yes, daily."
    assert sheet["D6"].value == NO_DATA_LABEL["en"]
    assert load_workbook(out)["Instructions"]["A1"].value.startswith("Fill in")


def test_failed_check_withholds_the_file_and_explains_why(db_conn, tmp_path, llm_settings, monkeypatch):
    path = security_workbook(tmp_path / "s.xlsx")
    job = create(db_conn, path, llm_settings, FakeLayoutLLM(SECURITY_MAP))
    db_conn.execute("UPDATE job_items SET status = 'answered', text = 'Yes', confidence = 0.9")
    db_conn.execute("UPDATE jobs SET status = 'done' WHERE id = %s", (job.id,))

    def broken(*args, **kwargs):
        raise ExportVerificationError(["sheet 'Security', cell B3: content changed"])

    monkeypatch.setattr(export_module, "verify_xlsx", broken)

    with pytest.raises(ConflictError) as caught:
        jobs.export_job(db_conn, job.id, "tester")
    assert "лист «Security», ячейка B3" in ru(str(caught.value))
