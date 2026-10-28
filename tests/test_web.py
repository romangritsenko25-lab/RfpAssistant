import time

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from rfp.auth import create_user
from tests.conftest import ADMIN_PASSWORD, ADMIN_USER, KB_TEXT

pytestmark = pytest.mark.db

XLSX_PARAMS = {
    "sheet": "Questionnaire",
    "question_col": "B",
    "answer_col": "C",
    "header_row": "2",
    "number_col": "A",
}


def login(client, username=ADMIN_USER, password=ADMIN_PASSWORD):
    return client.post("/login", data={"username": username, "password": password}, follow_redirects=False)


@pytest.fixture
def web(client):
    response = login(client)
    assert response.status_code == 303
    return client


def upload_kb(client):
    response = client.post(
        "/knowledge/upload",
        files={"file": ("backup.md", KB_TEXT.encode(), "text/markdown")},
        follow_redirects=False,
    )
    assert response.status_code == 303, response.text


def upload_job(client, path) -> int:
    with path.open("rb") as handle:
        response = client.post(
            "/jobs", files={"file": (path.name, handle)}, data=XLSX_PARAMS, follow_redirects=False
        )
    assert response.status_code == 303, response.text
    return int(response.headers["location"].rsplit("/", 1)[1])


def wait_done(client, job_id):
    deadline = time.time() + 20
    while time.time() < deadline:
        if client.get(f"/api/jobs/{job_id}/status").json()["status"] in ("done", "failed"):
            return
        time.sleep(0.05)
    raise AssertionError("job did not finish")


@pytest.fixture
def job_id(web, questionnaire_xlsx) -> int:
    upload_kb(web)
    job = upload_job(web, questionnaire_xlsx)
    wait_done(web, job)
    return job


# --- access ---------------------------------------------------------------------------------------


@pytest.mark.parametrize("path", ["/jobs", "/jobs/1", "/knowledge", "/logs"])
def test_pages_redirect_to_login_without_session(client, path):
    response = client.get(path, follow_redirects=False)
    assert response.status_code == 303
    assert response.headers["location"].startswith("/login")


def test_htmx_request_without_session_gets_client_side_redirect(client):
    response = client.get("/jobs/1/items/q1", headers={"HX-Request": "true"}, follow_redirects=False)
    assert response.headers["HX-Redirect"] == "/login"


def test_login_page_and_wrong_password(client):
    assert client.get("/login").status_code == 200
    response = login(client, password="wrong-password")
    assert response.status_code == 401
    assert "Неверный логин или пароль" in response.text
    assert "rfp_session" not in response.cookies


def test_login_returns_to_the_requested_page_and_ignores_foreign_targets(client):
    ok = client.post(
        "/login",
        data={"username": ADMIN_USER, "password": ADMIN_PASSWORD, "next": "/knowledge"},
        follow_redirects=False,
    )
    assert ok.headers["location"] == "/knowledge"
    client.cookies.clear()
    evil = client.post(
        "/login",
        data={"username": ADMIN_USER, "password": ADMIN_PASSWORD, "next": "//evil.example"},
        follow_redirects=False,
    )
    assert evil.headers["location"] == "/jobs"


def test_logout_ends_the_session(web):
    assert web.post("/logout", follow_redirects=False).headers["location"] == "/login"
    assert web.get("/jobs", follow_redirects=False).status_code == 303


def test_logs_are_for_administrators_only(web, db_conn):
    create_user(db_conn, "reviewer", "reviewer-pass-1")
    web.cookies.clear()
    assert login(web, "reviewer", "reviewer-pass-1").status_code == 303
    assert web.get("/jobs").status_code == 200
    forbidden = web.get("/logs")
    assert forbidden.status_code == 403
    assert "Журналы" not in forbidden.text.split("<main")[0]  # no logs link in the menu


# --- pages ----------------------------------------------------------------------------------------


def test_static_assets_are_served(client):
    assert client.get("/static/app.css").status_code == 200
    assert client.get("/static/htmx-2.0.4.min.js").status_code == 200
    assert client.get("/static/fonts/ibm-plex-sans-cyrillic-400.woff2").status_code == 200


def test_every_page_renders(web, job_id):
    for path in (
        "/jobs",
        f"/jobs/{job_id}",
        f"/jobs/{job_id}?f=review",
        f"/jobs/{job_id}/summary",
        f"/jobs/{job_id}/progress",
        "/jobs/table",
        "/knowledge",
        "/logs",
        "/logs?tab=model",
        f"/logs?tab=model&job_id={job_id}",
    ):
        response = web.get(path)
        assert response.status_code == 200, path
    assert web.get("/jobs/9999").status_code == 404


def test_job_list_shows_counts_and_review_table_shows_items(web, job_id):
    listing = web.get("/jobs").text
    assert "questionnaire.xlsx" in listing and "3 / 3" in listing
    page = web.get(f"/jobs/{job_id}").text
    assert "Do you enforce multi-factor authentication?" in page
    assert "Ответ готов" in page and "Нет данных" in page and "На проверку" in page
    assert "backup" in page  # the citation's document
    only_review = web.get(f"/jobs/{job_id}?f=review").text
    assert "Do you enforce multi-factor authentication?" not in only_review
    assert only_review.count('<tbody class="item') == 1  # q3 failed -> needs_review


def test_finished_job_progress_reloads_the_page(web, job_id):
    response = web.get(f"/jobs/{job_id}/progress")
    assert response.headers["HX-Refresh"] == "true"
    assert "hx-trigger" not in response.text  # polling stops


def test_bad_job_upload_shows_an_error_on_the_page(web):
    response = web.post("/jobs", files={"file": ("a.txt", b"x")})
    assert response.status_code == 422
    assert "Не удалось принять файл" in response.text


# --- review via HTMX --------------------------------------------------------------------------------


def test_edit_answer_returns_the_updated_row(web, job_id):
    form = web.get(f"/jobs/{job_id}/items/q1/edit", headers={"HX-Request": "true"})
    assert form.status_code == 200 and "<textarea" in form.text

    response = web.post(
        f"/jobs/{job_id}/items/q1/edit",
        data={"text": "We back up nightly at 02:00."},
        headers={"HX-Request": "true"},
    )
    assert response.status_code == 200
    assert response.headers["HX-Trigger"] == "item-changed"
    assert response.text.lstrip().startswith("<tbody")
    assert "We back up nightly at 02:00." in response.text
    assert "Изменён" in response.text
    assert "Backups run daily." in response.text  # the model's draft stays visible
    assert web.get(f"/api/jobs/{job_id}").json()["items"][0]["status"] == "edited"


def test_edit_with_empty_text_keeps_the_form_with_an_error(web, job_id):
    response = web.post(f"/jobs/{job_id}/items/q1/edit", data={"text": "   "})
    assert response.status_code == 200
    assert "<textarea" in response.text and 'role="alert"' in response.text


def test_approve_and_save_as_reference(web, job_id):
    refused = web.post(f"/jobs/{job_id}/items/q1/reference")
    assert 'role="alert"' in refused.text  # an unreviewed answer cannot become a reference
    assert "Сохранить как эталон" not in refused.text

    approved = web.post(f"/jobs/{job_id}/items/q2/approve")
    assert approved.headers["HX-Trigger"] == "item-changed"
    assert "Одобрен" in approved.text and "Сохранить как эталон" in approved.text

    saved = web.post(f"/jobs/{job_id}/items/q2/reference")
    assert "эталон REF-" in saved.text
    documents = web.get("/knowledge").text
    assert "REF-" in documents and "эталон" in documents


def test_summary_counts_follow_the_review(web, job_id):
    before = web.get(f"/jobs/{job_id}/summary", params={"f": "all"}).text
    web.post(f"/jobs/{job_id}/items/q3/approve")
    after = web.get(f"/jobs/{job_id}/summary", params={"f": "all"}).text
    assert before != after
    assert 'aria-current="true"' in after


def test_row_actions_on_missing_items_return_404(web, job_id):
    assert web.post(f"/jobs/{job_id}/items/zz/approve").status_code == 404
    assert web.get(f"/jobs/{job_id}/items/zz").status_code == 404


def test_export_downloads_the_file(web, job_id):
    response = web.get(f"/jobs/{job_id}/export")
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    assert response.content[:2] == b"PK"  # an xlsx is a zip
    assert web.get("/jobs/9999/export").status_code == 404


# --- knowledge base and logs -----------------------------------------------------------------------


def test_knowledge_base_upload_reindex_delete(web, api_settings):
    upload_kb(web)
    page = web.get("/knowledge").text
    assert "backup.md" in page
    assert (api_settings.kb_dir / "backup.md").exists()

    again = web.post(
        "/knowledge/upload", files={"file": ("backup.md", KB_TEXT.encode())}, follow_redirects=True
    )
    assert "не изменился" in again.text

    reindexed = web.post("/knowledge/reindex", follow_redirects=True)
    assert "Синхронизация завершена" in reindexed.text

    deleted = web.post("/knowledge/backup/delete", follow_redirects=True)
    assert "удалён из базы знаний" in deleted.text
    assert not (api_settings.kb_dir / "backup.md").exists()
    missing = web.post("/knowledge/backup/delete", follow_redirects=True)
    assert "не найден" in missing.text


def test_bad_knowledge_upload_shows_an_error(web):
    response = web.post("/knowledge/upload", files={"file": ("a.exe", b"x")})
    assert response.status_code == 422
    assert "Не удалось загрузить документ" in response.text


def test_logs_show_actions_and_model_requests(web, job_id):
    web.post(f"/jobs/{job_id}/items/q1/approve")
    actions = web.get("/logs").text
    for name in ("login", "job.create", "answer.edit", "document.upload"):
        assert name in actions
    requests = web.get("/logs?tab=model").text
    assert "gpt-4o-mini" in requests and "model unavailable" in requests
    assert "$0.0010" in requests
    assert "#" in requests and "backup" in requests  # fragment ids with their documents


def test_service_messages_are_shown_in_russian(web):
    from rfp.web.common import ru

    assert ru("2 model citation(s) rejected: not found in the fragments.").startswith(
        "Отклонено цитат модели: 2"
    )
    assert ru("Something unexpected") == "Something unexpected"
    response = web.post("/jobs", files={"file": ("a.txt", b"x")})
    assert "Неподдерживаемый тип файла" in response.text and "Unsupported" not in response.text


def test_brand_is_one_word(web):
    header = web.get("/jobs").text.split("</header>")[0]
    assert "Rfp<span>Assistant</span>" in header and "Rfp Assistant" not in header


# --- automatic column detection ---------------------------------------------------------------------

LAYOUT_QUESTIONS = [
    "Do you enforce multi-factor authentication?",
    "Как часто пересматриваются права доступа?",
    "Деректерді сақтау мерзімі қандай?",
]


def build_id_domain_question_answer(path):
    """ID / Domain / Question / Answer: the layout that the old B/C defaults destroyed."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Questionnaire"
    sheet.append(["ID", "Domain", "Question", "Answer"])
    for number, question in enumerate(LAYOUT_QUESTIONS, 1):
        sheet.append([f"Q-{number:03d}", "Security", question, None])
    sheet["A1"].font = Font(bold=True)
    sheet["C2"].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet.column_dimensions["C"].width = 60
    workbook.save(path)
    return path


def post_job(client, path, **data):
    with path.open("rb") as handle:
        return client.post("/jobs", files={"file": (path.name, handle)}, data=data, follow_redirects=False)


def test_upload_without_columns_detects_layout_and_export_keeps_the_questions(web, tmp_path):
    upload_kb(web)
    path = build_id_domain_question_answer(tmp_path / "layout.xlsx")
    response = post_job(web, path)
    assert response.status_code == 303, response.text
    job = int(response.headers["location"].rsplit("/", 1)[1])
    wait_done(web, job)

    page = web.get(f"/jobs/{job}").text
    assert "Лист: Questionnaire, вопросы: колонка C, ответы: колонка D" in page

    export = web.get(f"/jobs/{job}/export")
    assert export.status_code == 200
    out = tmp_path / "out.xlsx"
    out.write_bytes(export.content)
    original, result = load_workbook(path)["Questionnaire"], load_workbook(out)["Questionnaire"]
    assert [result[f"C{row}"].value for row in (2, 3, 4)] == LAYOUT_QUESTIONS
    assert all(result[f"D{row}"].value for row in (2, 3, 4))
    for row in range(1, 5):
        for col in "ABC":
            assert result[f"{col}{row}"].value == original[f"{col}{row}"].value
    assert result["A1"].font.bold
    assert result["C2"].fill.fgColor.rgb == original["C2"].fill.fgColor.rgb
    assert result.column_dimensions["C"].width == 60


def test_upload_with_equal_question_and_answer_columns_is_rejected(web, questionnaire_xlsx):
    response = post_job(web, questionnaire_xlsx, question_col="B", answer_col="b")
    assert response.status_code == 422
    assert "Колонка ответов должна отличаться" in response.text


def test_upload_without_recognisable_headers_asks_for_manual_columns(web, tmp_path):
    path = tmp_path / "plain.xlsx"
    workbook = Workbook()
    workbook.active.append(["Foo", "Bar"])
    workbook.active.append(["a", "b"])
    workbook.save(path)
    response = post_job(web, path)
    assert response.status_code == 422
    assert "Укажите колонки вручную" in response.text
