"""Several-part answers: the short rating, writing into the right columns, response-column rule, stubs."""

import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from rfp.answering import answer_question
from rfp.config import Settings
from rfp.export import ExportVerificationError, export_answers
from rfp.export.targets import match_allowed
from rfp.export.verify import verify_xlsx
from rfp.llm import LLMReply, LLMUsage
from rfp.models import Answer, Citation, Location, Question, RetrievedChunk
from rfp.questions.layout_map import AnswerTargets, SheetMap
from rfp.questions.layout_validate import check_sheet_map
from rfp.questions.placeholders import is_placeholder
from rfp.questions.xlsx_questions import extract_xlsx_questions
from rfp.questions.xlsx_snapshot import load_snapshot

CITATION = Citation(doc_id="DOC-1", section="Backups", snippet="Backups run daily", chunk_id=1)
YES_NO_NA = ["Yes", "No", "N/A"]


class FakeLLM:
    def __init__(self, **fields) -> None:
        base = {"text": "Yes [1].", "status": "answered", "confidence": 0.9}
        base["citations"] = [{"fragment": 1, "quote": "Backups of production databases are made daily"}]
        self.content = json.dumps({**base, **fields})

    def complete_json(self, system: str, user: str) -> LLMReply:
        return LLMReply(
            content=self.content, usage=LLMUsage(model="fake", prompt_tokens=1, completion_tokens=1)
        )


def generate(tmp_path: Path, **fields) -> Answer:
    prompt = tmp_path / "answer.md"
    prompt.write_text("P", encoding="utf-8")
    settings = Settings(
        database_url="postgresql://unused",
        answer_prompt_path=prompt,
        confidence_threshold=0.5,
        _env_file=None,
    )
    chunk = RetrievedChunk(
        chunk_id=1,
        doc_id="DOC-1",
        title="Policy",
        heading="Backups",
        text="Backups of production databases are made daily.",
        language="en",
        score=1.0,
    )
    question = Question(qid="q1", text="Are backups made?", location=Location())
    return answer_question(question, [chunk], FakeLLM(**fields), settings).answer


# --- the short rating ---------------------------------------------------------------------------


def test_short_rating_is_kept_for_answered(tmp_path: Path) -> None:
    assert generate(tmp_path, short=" Yes ").short == "yes"


@pytest.mark.parametrize("fields", [{"short": "maybe"}, {"short": 5}, {}])
def test_unknown_short_rating_is_dropped(tmp_path: Path, fields: dict) -> None:
    assert generate(tmp_path, **fields).short is None


def test_no_short_for_no_data_and_needs_review(tmp_path: Path) -> None:
    assert generate(tmp_path, status="no_data", short="no").short is None
    assert generate(tmp_path, status="needs_review", short="yes").short is None
    assert generate(tmp_path, confidence=0.1, short="yes").short is None  # demoted to needs_review


@pytest.mark.parametrize(
    ("short", "allowed", "expected"),
    [
        ("yes", ["Yes", "No", "N/A"], "Yes"),
        ("yes", ["да", "нет"], "да"),
        ("yes", ["Иә", "Жоқ"], "Иә"),
        ("no", ["  YES ", " no "], " no "),
        ("no", ["Да", "Нет"], "Нет"),
        ("n/a", ["Yes", "No", "Not Applicable"], "Not Applicable"),
        ("n/a", ["Yes", "No", "N/A"], "N/A"),
        ("partial", ["Yes - fully", "Partial", "No"], "Partial"),
        ("partial", ["Yes", "No"], None),
        ("yes", [], None),
        (None, ["Yes"], None),
    ],
)
def test_short_rating_is_matched_to_the_sheet_values(short, allowed, expected) -> None:
    assert match_allowed(short, allowed) == expected


# --- writing into columns -----------------------------------------------------------------------


def sheet_workbook(path: Path, marks: bool = False) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Q"
    if marks:
        sheet.append(["ID", "Question", "Yes", "No", "Notes"])
        sheet.append(["E-0", "Example row", "x", None, None])
    else:
        sheet.append(["ID", "Question", "Choice", "Notes", "Source"])
        validation = DataValidation(type="list", formula1='"Yes,No,N/A"')
        validation.add("C2:C9")
        sheet.add_data_validation(validation)
    sheet.append(["1", "Are backups made daily?"])
    sheet.append(["2", "Is there a DR site?"])
    workbook.save(path)
    return path


def questions_for(path: Path, targets: AnswerTargets, header_row: int = 1) -> list[Question]:
    return extract_xlsx_questions(path, "Q", "B", targets.text_col, header_row, "A", targets=targets)


def answers_for(questions: list[Question]) -> dict[str, Answer]:
    first, second = questions[-2:]
    return {
        first.qid: Answer(
            qid=first.qid, text="Daily.", status="answered", confidence=0.9, citations=[CITATION], short="yes"
        ),
        second.qid: Answer(qid=second.qid, text="", status="no_data", confidence=1.0, short="no"),
    }


def test_short_text_and_source_go_to_their_columns(tmp_path: Path) -> None:
    source = sheet_workbook(tmp_path / "s.xlsx")
    targets = AnswerTargets(short_answer_col="C", allowed_values=YES_NO_NA, text_col="D", source_col="E")
    questions = questions_for(source, targets)
    dest = tmp_path / "out.xlsx"

    export_answers(source, questions, answers_for(questions), dest)

    sheet = load_workbook(dest)["Q"]
    assert [sheet[f"{c}2"].value for c in "CDE"] == ["Yes", "Daily.", "DOC-1, Backups"]
    assert sheet["C3"].value is None  # no_data: no short answer
    assert sheet["D3"].value == "No data in the knowledge base"
    assert sheet["E3"].value is None


def test_without_source_column_sources_stay_in_the_text(tmp_path: Path) -> None:
    source = sheet_workbook(tmp_path / "s.xlsx")
    targets = AnswerTargets(short_answer_col="C", allowed_values=YES_NO_NA, text_col="D")
    questions = questions_for(source, targets)
    dest = tmp_path / "out.xlsx"

    export_answers(source, questions, answers_for(questions), dest)

    assert load_workbook(dest)["Q"]["D2"].value == "Daily.\nSources: DOC-1, Backups"


def test_value_outside_the_list_is_not_written(tmp_path: Path) -> None:
    source = sheet_workbook(tmp_path / "s.xlsx")
    targets = AnswerTargets(short_answer_col="C", allowed_values=["Yes", "No"], text_col="D")
    questions = questions_for(source, targets)
    answers = answers_for(questions)
    answers[questions[0].qid] = answers[questions[0].qid].model_copy(update={"short": "partial"})
    dest = tmp_path / "out.xlsx"

    export_answers(source, questions, answers, dest)

    assert load_workbook(dest)["Q"]["C2"].value is None


def test_self_check_rejects_a_short_cell_outside_the_list(tmp_path: Path) -> None:
    source = sheet_workbook(tmp_path / "s.xlsx")
    targets = AnswerTargets(short_answer_col="C", allowed_values=YES_NO_NA, text_col="D", source_col="E")
    questions = questions_for(source, targets)
    answers = answers_for(questions)
    dest = tmp_path / "out.xlsx"
    export_answers(source, questions, answers, dest)
    workbook = load_workbook(dest)
    workbook["Q"]["C2"] = "Maybe"
    workbook.save(dest)

    with pytest.raises(ExportVerificationError, match="not an allowed value"):
        verify_xlsx(source, dest, questions, answers)


def test_mark_goes_to_the_chosen_column_with_the_sign_of_the_sheet(tmp_path: Path) -> None:
    source = sheet_workbook(tmp_path / "m.xlsx", marks=True)
    targets = AnswerTargets(text_col="E", choice_cols={"Yes": "C", "No": "D"})
    questions = questions_for(source, targets)
    dest = tmp_path / "out.xlsx"

    export_answers(source, questions, answers_for(questions), dest)

    sheet = load_workbook(dest)["Q"]
    assert (sheet["C3"].value, sheet["D3"].value) == ("x", None)  # the sample row used "x"
    assert (sheet["C4"].value, sheet["D4"].value) == (None, None)  # no_data: no mark


def test_mark_defaults_to_a_cross(tmp_path: Path) -> None:
    source = sheet_workbook(tmp_path / "m.xlsx", marks=True)
    workbook = load_workbook(source)
    workbook["Q"]["C2"] = None
    workbook.save(source)
    targets = AnswerTargets(text_col="E", choice_cols={"Yes": "C", "No": "D"})
    questions = questions_for(source, targets)
    dest = tmp_path / "out.xlsx"

    export_answers(source, questions, answers_for(questions), dest)

    assert load_workbook(dest)["Q"]["C3"].value == "X"


# --- the response column rule -------------------------------------------------------------------


def comments_workbook(path: Path, filled: bool = False) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Q"
    sheet.append(["ID", "Question", "Yes/No", "Comments", "Owner"])
    for n in range(1, 5):
        sheet.append([str(n), f"Question {n} text?", None, "see policy" if filled else None, None])
    workbook.save(path)
    return path


def add_list(path: Path, cell_range: str) -> None:
    workbook = load_workbook(path)
    validation = DataValidation(type="list", formula1='"Yes,No"')
    validation.add(cell_range)
    workbook["Q"].add_data_validation(validation)
    workbook.save(path)


def model_map(**targets) -> SheetMap:
    return SheetMap.model_validate(
        {
            "sheet": "Q",
            "header_row": 1,
            "question_col": "B",
            "number_col": "A",
            "answer_targets": targets,
        }
    )


def test_model_text_column_is_moved_to_the_empty_comments_column(tmp_path: Path) -> None:
    snap = load_snapshot(comments_workbook(tmp_path / "c.xlsx"))["Q"]

    checked = check_sheet_map(model_map(text_col="F"), snap)

    assert checked.answer_targets.text_col == "D"
    assert checked.answer_header is None


def test_filled_comments_column_is_not_taken(tmp_path: Path) -> None:
    snap = load_snapshot(comments_workbook(tmp_path / "c.xlsx", filled=True))["Q"]

    assert check_sheet_map(model_map(text_col="F"), snap).answer_targets.text_col == "F"


def test_drop_down_column_titled_answer_is_not_a_text_candidate(tmp_path: Path) -> None:
    path = comments_workbook(tmp_path / "c.xlsx")
    workbook = load_workbook(path)
    workbook["Q"]["C1"] = "Answer"
    workbook["Q"]["D1"] = "Owner"
    workbook.save(path)
    add_list(path, "C2:C5")

    checked = check_sheet_map(model_map(text_col="F"), load_snapshot(path)["Q"])

    assert checked.answer_targets.text_col == "F"


def test_list_in_the_sheet_becomes_the_allowed_values(tmp_path: Path) -> None:
    path = comments_workbook(tmp_path / "c.xlsx")
    add_list(path, "C2:C5")

    checked = check_sheet_map(model_map(short_answer_col="C", text_col="D"), load_snapshot(path)["Q"])

    assert checked.answer_targets.allowed_values == ["No", "Yes"]


# --- template stubs in DOCX ---------------------------------------------------------------------


@pytest.mark.parametrize(
    "text",
    [
        "[insert response]",
        "Click here to enter text.",
        "[Vendor to provide answer]",
        "",
        "_____",
        "Нажмите здесь, чтобы ввести текст",
    ],
)
def test_stubs_count_as_empty(text: str) -> None:
    assert is_placeholder(text)


@pytest.mark.parametrize("text", ["We encrypt data at rest.", "[1] Policy", "N/A", "Click Save to continue"])
def test_real_text_is_not_a_stub(text: str) -> None:
    assert not is_placeholder(text)


def test_stub_cell_is_replaced_by_the_answer(tmp_path: Path) -> None:
    source = tmp_path / "t.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Question"
    table.cell(0, 1).text = "Response"
    table.cell(1, 0).text = "Do you encrypt data?"
    table.cell(1, 1).text = "[insert response]"
    doc.save(source)
    question = Question(
        qid="q1",
        text="Do you encrypt data?",
        location=Location(block=0, table_row=1, table_col=0, answer_table_row=1, answer_table_col=1),
    )
    answer = Answer(qid="q1", text="Yes, AES-256.", status="answered", confidence=0.9)
    dest = tmp_path / "out.docx"

    export_answers(source, [question], {"q1": answer}, dest)

    assert [p.text for p in Document(str(dest)).tables[0].cell(1, 1).paragraphs] == ["Yes, AES-256."]
