"""The model's layout map: how it is requested, checked by code and replaced by header rules."""

import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from rfp.llm import LLMReply, LLMUsage
from rfp.questions import ManualLayout, resolve_docx, resolve_xlsx
from rfp.questions.layout_map import SheetMap
from rfp.questions.layout_validate import LayoutRejected, check_sheet_map
from rfp.questions.xlsx_snapshot import load_snapshot, render_sheet
from rfp.services.errors import InvalidInputError

PROMPT = "LAYOUT PROMPT"


class FakeLayoutLLM:
    """Returns a prepared reply (a dict, a raw string or an exception) and remembers the requests."""

    def __init__(self, reply) -> None:
        self.reply = reply
        self.requests: list[tuple[str, str]] = []

    def complete_json(self, system: str, user: str) -> LLMReply:
        self.requests.append((system, user))
        if isinstance(self.reply, Exception):
            raise self.reply
        content = self.reply if isinstance(self.reply, str) else json.dumps(self.reply)
        return LLMReply(content=content, usage=LLMUsage(model="fake", prompt_tokens=10, completion_tokens=5))


def security_workbook(path: Path) -> Path:
    """Instructions sheet + a sheet with unconventional headers, section rows and a drop-down list."""
    workbook = Workbook()
    workbook.active.title = "Instructions"
    workbook["Instructions"]["A1"] = "Fill in the Security sheet. Do not change the questions."
    sheet = workbook.create_sheet("Security")
    sheet.append(["Ref", "Item text", "Choice", "Comment"])
    sheet.append([None, "Access control"])
    sheet.append(["AC-1", "Do you enforce MFA?"])
    sheet.append(["AC-2", "How often are access rights reviewed?"])
    sheet["B5"] = "Data protection"
    sheet.merge_cells("B5:D5")
    sheet.append(["DP-1", "Where is the data stored?"])
    validation = DataValidation(type="list", formula1='"Yes,No,Partial"')
    validation.add("C3:C6")
    sheet.add_data_validation(validation)
    workbook.save(path)
    return path


SECURITY_MAP = {
    "sheets": [
        {"sheet": "Instructions", "has_questions": False},
        {
            "sheet": "Security",
            "has_questions": True,
            "header_row": 1,
            "question_col": "B",
            "number_col": "A",
            "section_header_rows": [2],
            "first_question_row": 2,
            "last_question_row": 6,
            "answer_targets": {
                "short_answer_col": "C",
                "allowed_values": ["Yes", "No"],
                "text_col": "D",
                "source_col": None,
            },
        },
    ]
}


def plain_workbook(path: Path, header: list[str] | None = None) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(header or ["ID", "Question", "Answer"])
    sheet.append(["1", "Do you encrypt backups?", None])
    sheet.append(["2", "Who owns the data?", None])
    workbook.save(path)
    return path


def one_sheet(question_col: str, text_col: str) -> dict:
    sheet = {"sheet": "Sheet1", "header_row": 1, "question_col": question_col}
    return {"sheets": [{**sheet, "answer_targets": {"text_col": text_col}}]}


def variant(**changes) -> dict:
    """SECURITY_MAP with changes applied to the Security sheet's entry."""
    data = json.loads(json.dumps(SECURITY_MAP))
    targets = changes.pop("answer_targets", None)
    data["sheets"][1].update(changes)
    if targets:
        data["sheets"][1]["answer_targets"].update(targets)
    return data


def resolve(path: Path, reply, manual: ManualLayout | None = None):
    llm = FakeLayoutLLM(reply)
    return resolve_xlsx(path, manual or ManualLayout(), llm=llm, prompt=PROMPT), llm


# --- a valid map -------------------------------------------------------------------------------


def test_valid_map_is_applied_across_sheets(tmp_path):
    result, llm = resolve(security_workbook(tmp_path / "s.xlsx"), SECURITY_MAP)

    assert result.layout.source == "llm"
    assert [(q.qid, q.text, q.number) for q in result.questions] == [
        ("q1", "Do you enforce MFA?", "AC-1"),
        ("q2", "How often are access rights reviewed?", "AC-2"),
        ("q3", "Where is the data stored?", "DP-1"),
    ]
    assert {(q.location.sheet, q.location.answer_col, q.location.question_col) for q in result.questions} == {
        ("Security", "D", "B")
    }
    instructions, security = result.layout.sheets
    assert not instructions.has_questions
    assert security.skipped_rows == 2  # the listed title and the merged banner
    assert security.answer_targets.short_answer_col == "C"
    assert security.answer_targets.allowed_values == ["Yes", "No"]
    assert security.answer_header is None  # D1 already has a title
    assert len(llm.requests) == 1
    assert llm.requests[0][0] == PROMPT
    assert result.calls[0].usage.model == "fake"


def test_request_shows_addresses_merged_cells_and_dropdown_lists(tmp_path):
    _, llm = resolve(security_workbook(tmp_path / "s.xlsx"), SECURITY_MAP)

    request = llm.requests[0][1]
    assert request.startswith("FORMAT: XLSX")
    assert 'B3="Do you enforce MFA?"' in request
    assert "MERGED CELLS: B5:D5" in request
    assert "DROP-DOWN LIST at C3:C6: Yes, No, Partial" in request
    assert 'SHEET "Instructions"' in request


def test_dropdown_list_from_a_cell_range_is_read(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Q"
    lists = workbook.create_sheet("Lists")
    for row, value in enumerate(["High", "Low"], 1):
        lists.cell(row=row, column=1, value=value)
    sheet.append(["Question", "Level"])
    sheet.append(["How bad is it?"])
    validation = DataValidation(type="list", formula1="Lists!$A$1:$A$2")
    validation.add("B2:B9")
    sheet.add_data_validation(validation)
    workbook.save(tmp_path / "r.xlsx")

    snapshot = load_snapshot(tmp_path / "r.xlsx")["Q"]

    assert snapshot.list_values(2, range(2, 10)) == {"High", "Low"}
    assert "DROP-DOWN LIST at B2:B9: High, Low" in render_sheet(snapshot)


def test_long_sheet_is_shown_by_head_middle_and_tail(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Question"])
    for number in range(1, 200):
        sheet.append([f"Question number {number}?"])
    workbook.save(tmp_path / "long.xlsx")

    text = render_sheet(load_snapshot(tmp_path / "long.xlsx")["Sheet"])

    assert 'A2="Question number 1?"' in text and 'A25="Question number 24?"' in text
    assert 'A200="Question number 199?"' in text
    assert "not shown" in text and "A26=" not in text
    assert len(text.splitlines()) < 50


# --- rejected maps fall back to header rules ----------------------------------------------------


def test_answer_column_equal_to_question_column_falls_back_to_rules(tmp_path):
    reply = {
        "sheets": [
            {
                "sheet": "Sheet1",
                "header_row": 1,
                "question_col": "B",
                "answer_targets": {"text_col": "B"},
            }
        ]
    }
    result, _ = resolve(plain_workbook(tmp_path / "p.xlsx"), reply)

    assert result.layout.source == "rules"
    assert "also the question column" in result.layout.note
    assert [q.location.answer_col for q in result.questions] == ["C", "C"]
    assert result.calls[0].error  # the rejection is kept in the request log


def test_rows_outside_the_sheet_fall_back_to_rules(tmp_path):
    reply = {
        "sheets": [
            {
                "sheet": "Sheet1",
                "header_row": 1,
                "question_col": "B",
                "first_question_row": 50,
                "last_question_row": 60,
                "answer_targets": {"text_col": "C"},
            }
        ]
    }
    result, _ = resolve(plain_workbook(tmp_path / "p.xlsx"), reply)

    assert result.layout.source == "rules"
    assert "below the last row" in result.layout.note
    assert len(result.questions) == 2


def test_question_column_without_text_is_rejected(tmp_path):
    reply = {
        "sheets": [
            {"sheet": "Sheet1", "header_row": 1, "question_col": "D", "answer_targets": {"text_col": "C"}}
        ]
    }
    result, _ = resolve(plain_workbook(tmp_path / "p.xlsx"), reply)

    assert result.layout.source == "rules"
    assert "no question text in column D" in result.layout.note


def test_answer_column_that_holds_the_questions_is_rejected(tmp_path):
    path = plain_workbook(tmp_path / "p.xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["ID", "Question", "Answer"])
    sheet.append(["1", "Do you encrypt backups?", "Do you encrypt backups? yes"])
    sheet.append(["2", "Who owns the data?", "Who owns the data?"])
    workbook.save(path)
    reply = {
        "sheets": [
            {"sheet": "Sheet1", "header_row": 1, "question_col": "B", "answer_targets": {"text_col": "C"}}
        ]
    }

    result, _ = resolve(path, reply)

    assert result.layout.source == "rules"
    assert "holds the question text" in result.layout.note


def test_allowed_values_outside_the_dropdown_list_are_rejected(tmp_path):
    reply = variant(answer_targets={"allowed_values": ["Yes", "Maybe"]})
    path = security_workbook(tmp_path / "s.xlsx")

    with pytest.raises(InvalidInputError, match="Specify the columns manually"):
        resolve(path, reply)  # headers are unconventional, so the rules cannot help either


def test_allowed_values_without_a_dropdown_list_are_dropped(tmp_path):
    # no short-answer column means no list to check the values against: they are not kept
    reply = variant(answer_targets={"short_answer_col": None, "allowed_values": ["Yes"]})
    result, _ = resolve(security_workbook(tmp_path / "s.xlsx"), reply)

    assert result.layout.sheets[1].answer_targets.allowed_values is None


@pytest.mark.parametrize(
    "reply",
    [
        "this is not json",
        "[1, 2]",
        {"sheets": "nope"},
        {"sheets": [{"sheet": "Sheet1", "question_col": "1; DROP", "answer_targets": {"text_col": "C"}}]},
        {"sheets": [{"sheet": "Ghost", "question_col": "B", "answer_targets": {"text_col": "C"}}]},
        {"sheets": [{"sheet": "Sheet1", "has_questions": False}]},
        RuntimeError("model unavailable"),
    ],
)
def test_unusable_replies_fall_back_to_rules(tmp_path, reply):
    result, _ = resolve(plain_workbook(tmp_path / "p.xlsx"), reply)

    assert result.layout.source == "rules"
    assert result.layout.note
    assert len(result.questions) == 2


def test_no_valid_map_and_no_header_asks_for_manual_columns(tmp_path):
    path = security_workbook(tmp_path / "s.xlsx")

    with pytest.raises(InvalidInputError, match="Specify the columns manually"):
        resolve(path, "garbage")


# --- authority order ----------------------------------------------------------------------------


def test_manual_columns_win_and_the_model_is_not_asked(tmp_path):
    manual = ManualLayout(question_col="B", answer_col="C")
    result, llm = resolve(plain_workbook(tmp_path / "p.xlsx"), SECURITY_MAP, manual)

    assert result.layout.source == "manual"
    assert llm.requests == []
    assert result.layout.sheets[0].answer_targets.text_col == "C"


def test_without_a_model_header_rules_are_used(tmp_path):
    result = resolve_xlsx(plain_workbook(tmp_path / "p.xlsx"), ManualLayout(), llm=None, prompt=None)

    assert result.layout.source == "rules"
    assert result.layout.note is None and result.calls == []


def test_rules_look_past_an_instruction_sheet(tmp_path):
    workbook = Workbook()
    workbook.active.title = "Read me"
    workbook.active["A1"] = "Instructions"
    sheet = workbook.create_sheet("Form")
    sheet.append(["№", "Вопрос", "Ответ"])
    sheet.append(["1", "Как хранятся данные?", None])
    workbook.save(tmp_path / "f.xlsx")

    result = resolve_xlsx(tmp_path / "f.xlsx", ManualLayout(), llm=None, prompt=None)

    assert (result.layout.sheets[0].sheet, len(result.questions)) == ("Form", 1)


def test_a_new_answer_column_gets_a_title(tmp_path):
    reply = {
        "sheets": [
            {"sheet": "Sheet1", "header_row": 1, "question_col": "B", "answer_targets": {"text_col": "E"}}
        ]
    }
    result, _ = resolve(plain_workbook(tmp_path / "p.xlsx", ["ID", "Question", "Owner"]), reply)

    assert result.layout.source == "llm"
    assert result.layout.sheets[0].answer_header == "Answer"


def test_a_requested_sheet_limits_the_map(tmp_path):
    path = security_workbook(tmp_path / "s.xlsx")
    llm = FakeLayoutLLM(SECURITY_MAP)

    result = resolve_xlsx(path, ManualLayout(sheet="Security"), llm=llm, prompt=PROMPT)

    assert "Instructions" not in llm.requests[0][1]
    assert len(result.questions) == 3
    with pytest.raises(InvalidInputError, match="Sheet 'Nope' not found"):
        resolve_xlsx(path, ManualLayout(sheet="Nope"), llm=llm, prompt=PROMPT)


def test_check_refuses_an_answer_column_that_is_mostly_filled(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "S"
    sheet.append(["Question", "Notes"])
    for number in range(4):
        sheet.append([f"Question {number}?", "already written"])
    workbook.save(tmp_path / "f.xlsx")
    snapshot = load_snapshot(tmp_path / "f.xlsx")["S"]
    sheet_map = SheetMap.model_validate(
        {"sheet": "S", "header_row": 1, "question_col": "A", "answer_targets": {"text_col": "B"}}
    )

    with pytest.raises(LayoutRejected, match="already holds data"):
        check_sheet_map(sheet_map, snapshot)


# --- DOCX ---------------------------------------------------------------------------------------


def docx_form(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Please describe your backup policy.")
    table = doc.add_table(rows=3, cols=3)
    for row, values in zip(
        table.rows,
        [
            ("No", "Requirement", "Response"),
            ("1", "Does it support SAML?", ""),
            ("2", "Name the log retention period.", ""),
        ],
        strict=True,
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    doc.save(path)
    return path


def cell_item(row: int, answer_row: int, answer_col: int) -> dict:
    """A question in table block 1, column 1, answered in another cell of the same table."""
    place = {"place": "cell", "table_row": answer_row, "table_col": answer_col}
    return {"block": 1, "table_row": row, "table_col": 1, "answer": place}


DOCX_MAP = {
    "questions": [
        {"block": 0, "answer": {"place": "paragraph_below"}},
        {
            "block": 1,
            "table_row": 1,
            "table_col": 1,
            "answer": {"place": "cell", "table_row": 1, "table_col": 2},
        },
        {
            "block": 1,
            "table_row": 2,
            "table_col": 1,
            "answer": {"place": "cell", "table_row": 2, "table_col": 2},
        },
    ]
}


def test_docx_map_gives_question_and_answer_places(tmp_path):
    llm = FakeLayoutLLM(DOCX_MAP)

    result = resolve_docx(docx_form(tmp_path / "f.docx"), llm=llm, prompt=PROMPT)

    assert result.layout.source == "llm"
    assert (result.layout.docx_paragraph_answers, result.layout.docx_cell_answers) == (1, 2)
    first, second, third = result.questions
    assert (first.qid, first.text, first.location.answer_table_row) == (
        "q1",
        "Please describe your backup policy.",
        None,
    )
    assert (second.text, second.number) == ("Does it support SAML?", "1")
    assert (second.location.answer_table_row, second.location.answer_table_col) == (1, 2)
    assert third.location.answer_table_row == 2
    assert llm.requests[0][1].startswith("FORMAT: DOCX")
    assert "[1] TABLE 3 rows" in llm.requests[0][1]


@pytest.mark.parametrize(
    "item",
    [
        {"block": 9},
        {
            "block": 1,
            "table_row": 1,
            "table_col": 1,
            "answer": {"place": "cell", "table_row": 1, "table_col": 1},
        },
        {
            "block": 1,
            "table_row": 0,
            "table_col": 1,
            "answer": {"place": "cell", "table_row": 0, "table_col": 2},
        },
        {"block": 1, "table_row": 5, "table_col": 1},
        {
            "block": 1,
            "table_row": 1,
            "table_col": 1,
            "answer": {"place": "cell", "table_row": 1, "table_col": 7},
        },
        {"block": 0, "answer": {"place": "cell", "table_row": 1, "table_col": 2}},
    ],
)
def test_docx_map_breaking_a_rule_falls_back_to_rules(tmp_path, item):
    result = resolve_docx(
        docx_form(tmp_path / "f.docx"), llm=FakeLayoutLLM({"questions": [item]}), prompt=PROMPT
    )

    assert result.layout.source == "rules"
    assert result.layout.note
    assert len(result.questions) >= 1
