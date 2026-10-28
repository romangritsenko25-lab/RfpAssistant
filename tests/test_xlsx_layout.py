from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from rfp.export import AnswerHeader, export_answers
from rfp.models import Answer, Location, Question
from rfp.questions import detect_xlsx_layout, extract_xlsx_questions
from rfp.services.errors import InvalidInputError


def build(path: Path, rows: list[list[str | None]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def one_answer() -> dict[str, Answer]:
    return {"q1": Answer(qid="q1", text="Yes", status="answered", confidence=0.9, citations=[])}


@pytest.mark.parametrize(
    ("header", "expected"),
    [
        (["ID", "Domain", "Question", "Answer"], ("C", "D", "A")),
        (["№", "Раздел", "Вопрос", "Ответ поставщика"], ("C", "D", "A")),
        (["Нөмір", "Бөлім", "Сұрақ", "Жауап"], ("C", "D", "A")),
        ([" requirement ", "Vendor  Response:"], ("A", "B", None)),
        (["Талап", None, "Жауап"], ("A", "C", None)),
    ],
)
def test_headers_in_three_languages_are_detected(tmp_path, header, expected):
    path = build(tmp_path / "q.xlsx", [header, ["x"] * len(header)])
    layout = detect_xlsx_layout(path)
    assert (layout.question_col, layout.answer_col, layout.number_col) == expected
    assert (layout.header_row, layout.answer_header) == (1, None)


def test_header_is_found_below_a_title_row(tmp_path):
    path = build(tmp_path / "q.xlsx", [["RFP 2026"], [], ["ID", "Question", "Answer"]])
    layout = detect_xlsx_layout(path)
    assert (layout.header_row, layout.question_col, layout.answer_col) == (3, "B", "C")


@pytest.mark.parametrize(
    ("header", "title"),
    [
        (["ID", "Domain", "Question"], "Answer"),
        (["№", "Вопрос"], "Ответ"),
        (["Сұрақ"], "Жауап"),
    ],
)
def test_missing_answer_column_goes_right_of_the_headers_with_a_title(tmp_path, header, title):
    path = build(tmp_path / "q.xlsx", [header])
    layout = detect_xlsx_layout(path)
    assert layout.answer_col == "ABCDE"[len(header)]
    assert layout.answer_header == title


def test_explicit_values_win_over_detection(tmp_path):
    path = build(tmp_path / "q.xlsx", [["ID", "Domain", "Question", "Answer"]])
    layout = detect_xlsx_layout(path, question_col="b", answer_col="e", header_row=1)
    assert (layout.question_col, layout.answer_col) == ("B", "E")


def test_no_recognisable_headers_is_a_clear_error(tmp_path):
    path = build(tmp_path / "q.xlsx", [["Foo", "Bar"], ["a", "b"]])
    with pytest.raises(InvalidInputError, match="Specify the columns manually"):
        detect_xlsx_layout(path)


def test_same_question_and_answer_column_is_refused(tmp_path):
    path = build(tmp_path / "q.xlsx", [["ID", "Question", "Answer"]])
    with pytest.raises(InvalidInputError, match="must differ"):
        detect_xlsx_layout(path, question_col="B", answer_col="b")


def test_export_never_overwrites_the_question_column(tmp_path):
    path = build(tmp_path / "q.xlsx", [["ID", "Question"], ["1", "Is it safe?"]])
    questions = extract_xlsx_questions(path, "Sheet1", "B", "B")
    with pytest.raises(ValueError, match="question column"):
        export_answers(path, questions, one_answer(), tmp_path / "out.xlsx")


def test_export_refuses_a_cell_that_holds_the_question_text(tmp_path):
    path = build(tmp_path / "q.xlsx", [["ID", "Question", "Answer"], ["1", "Is it safe?", "Is it safe?"]])
    question = Question(
        qid="q1", text="Is it safe?", location=Location(sheet="Sheet1", row=2, answer_col="C")
    )
    with pytest.raises(ValueError, match="already holds the question"):
        export_answers(path, [question], one_answer(), tmp_path / "out.xlsx")


def test_export_writes_the_title_of_an_added_answer_column(tmp_path):
    path = build(tmp_path / "q.xlsx", [["ID", "Question"], ["1", "Is it safe?"]])
    workbook = load_workbook(path)
    workbook.active["B1"].font = Font(bold=True)
    workbook.save(path)
    question = Question(
        qid="q1",
        text="Is it safe?",
        location=Location(sheet="Sheet1", row=2, answer_col="C", question_col="B"),
    )
    out = tmp_path / "out.xlsx"
    export_answers(path, [question], one_answer(), out, AnswerHeader("Sheet1", 1, "C", "Answer"))
    sheet = load_workbook(out)["Sheet1"]
    assert sheet["C1"].value == "Answer"
    assert sheet["C1"].font.bold
    assert sheet["C2"].value.startswith("Yes")
    assert sheet["B2"].value == "Is it safe?"
