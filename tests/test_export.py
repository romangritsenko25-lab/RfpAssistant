from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from rfp.export import export_answers
from rfp.export.docx_export import ANSWER_STYLE
from rfp.models import Answer, Citation, Question
from rfp.questions import extract_docx_questions, extract_xlsx_questions
from tests import factories


def answers_for(questions: list[Question]) -> dict[str, Answer]:
    return {
        q.qid: Answer(
            qid=q.qid,
            text=f"Answer to {q.qid}",
            status="answered",
            confidence=0.9,
            citations=[Citation(doc_id="policy", section="Access", snippet="...")],
        )
        for q in questions
    }


def test_xlsx_answers_land_in_their_cells_and_styles_survive(
    questionnaire_xlsx: Path, tmp_path: Path
) -> None:
    questions = extract_xlsx_questions(questionnaire_xlsx, "Questionnaire", "B", "C", header_row=2)
    answers = answers_for(questions[:2])  # the third question stays unanswered
    dest = tmp_path / "out.xlsx"

    export_answers(questionnaire_xlsx, questions, answers, dest)

    sheet = load_workbook(dest)["Questionnaire"]
    assert sheet["C3"].value == "Answer to q1\nSources: policy, Access"
    assert sheet["C4"].value == "Answer to q2\nSources: policy, Access"
    assert sheet["C6"].value is None
    assert sheet["B3"].value == "Do you enforce multi-factor authentication?"


def test_xlsx_export_preserves_structure_and_formatting(questionnaire_xlsx: Path, tmp_path: Path) -> None:
    questions = extract_xlsx_questions(questionnaire_xlsx, "Questionnaire", "B", "C", header_row=2)
    dest = tmp_path / "out.xlsx"

    export_answers(questionnaire_xlsx, questions, answers_for(questions), dest)

    before, after = load_workbook(questionnaire_xlsx), load_workbook(dest)
    assert after.sheetnames == before.sheetnames
    assert after["Notes"]["A1"].value == "Keep me"
    sheet = after["Questionnaire"]
    assert (
        sheet["B2"].font.bold
        and sheet["B2"].fill.fgColor.rgb == before["Questionnaire"]["B2"].fill.fgColor.rgb
    )
    assert sheet["C3"].fill.fgColor.rgb == before["Questionnaire"]["C3"].fill.fgColor.rgb
    assert (sheet.column_dimensions["B"].width, sheet.column_dimensions["C"].width) == (60, 45)
    changed = {
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if cell.value != before["Questionnaire"][cell.coordinate].value
    }
    assert changed == {"C3", "C4", "C6"}


def test_xlsx_export_leaves_source_file_untouched(questionnaire_xlsx: Path, tmp_path: Path) -> None:
    questions = extract_xlsx_questions(questionnaire_xlsx, "Questionnaire", "B", "C", header_row=2)
    original = questionnaire_xlsx.read_bytes()

    export_answers(questionnaire_xlsx, questions, answers_for(questions), tmp_path / "out.xlsx")

    assert questionnaire_xlsx.read_bytes() == original


def body_texts(path: Path) -> list[str]:
    return [p.text for p in Document(str(path)).paragraphs]


def test_docx_answer_follows_each_question(rfp_docx: Path, tmp_path: Path) -> None:
    questions = extract_docx_questions(rfp_docx)
    answers = answers_for(questions)
    dest = tmp_path / "out.docx"

    export_answers(rfp_docx, questions, answers, dest)

    texts = body_texts(dest)
    for question in questions:
        if question.location.table_row is not None:
            continue
        position = next(i for i, text in enumerate(texts) if question.text in text)
        assert texts[position + 1].startswith(answers[question.qid].text)


def test_docx_export_changes_nothing_but_inserted_answers(rfp_docx: Path, tmp_path: Path) -> None:
    questions = extract_docx_questions(rfp_docx)
    dest = tmp_path / "out.docx"

    export_answers(rfp_docx, questions, answers_for(questions), dest)

    result = Document(str(dest))
    is_answer = [p.style is not None and p.style.name == ANSWER_STYLE for p in result.paragraphs]
    kept = [p.text for p, answer in zip(result.paragraphs, is_answer, strict=True) if not answer]
    assert kept == body_texts(rfp_docx)
    assert sum(is_answer) == sum(q.location.table_row is None for q in questions)
    assert [len(t.rows) for t in result.tables] == [len(t.rows) for t in Document(str(rfp_docx)).tables]


def test_docx_table_answer_goes_into_the_question_cell(rfp_docx: Path, tmp_path: Path) -> None:
    questions = extract_docx_questions(rfp_docx)
    dest = tmp_path / "out.docx"

    export_answers(rfp_docx, questions, answers_for(questions), dest)

    saml = next(q for q in questions if "SAML" in q.text)
    cell = Document(str(dest)).tables[0].cell(saml.location.table_row, saml.location.table_col)
    assert [p.text for p in cell.paragraphs][-1].startswith(f"Answer to {saml.qid}")


def test_docx_sources_are_written_in_small_text(rfp_docx: Path, tmp_path: Path) -> None:
    questions = extract_docx_questions(rfp_docx)[:1]
    dest = tmp_path / "out.docx"

    export_answers(rfp_docx, questions, answers_for(questions), dest)

    answer = next(p for p in Document(str(dest)).paragraphs if p.style.name == ANSWER_STYLE)
    assert answer.runs[-1].text.endswith("Sources: policy, Access")
    assert answer.runs[-1].font.size.pt == 8


def test_pdf_input_produces_docx_report(tmp_path: Path) -> None:
    pdf = factories.build_pdf(tmp_path / "rfp.pdf", [["1. Describe your backup policy."]])
    question = Question(
        qid="q1", text="Describe your backup policy.", number="1", location={"page": 1, "line": 1}
    )
    answers = {"q1": Answer(qid="q1", text="Nightly.", status="answered", confidence=0.9)}
    dest = tmp_path / "report.docx"

    export_answers(pdf, [question], answers, dest)

    texts = body_texts(dest)
    assert texts[-2:] == ["1. Describe your backup policy.", "Nightly."]


def test_unsupported_source_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="Unsupported"):
        export_answers(tmp_path / "x.txt", [], {}, tmp_path / "out.txt")
