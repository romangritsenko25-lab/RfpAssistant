"""Hard cases found on real questionnaires, rebuilt as small synthetic workbooks."""

import json
import re
import zipfile
from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import Workbook

from rfp.export.verify import verify_xlsx
from rfp.export.xlsx_export import export_xlsx
from rfp.llm import LLMReply, LLMUsage
from rfp.models import Answer, Location, Question
from rfp.questions import ManualLayout, resolve_xlsx
from rfp.questions.layout_map import SheetMap
from rfp.questions.layout_repeats import find_repeat_sheets
from rfp.questions.layout_validate import LayoutRejected, check_sheet_map
from rfp.questions.xlsx_snapshot import load_snapshot
from rfp.questions.xlsx_x14 import read_x14_lists

X14 = (
    '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
    'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    '<x14:dataValidations xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main" count="1">'
    '<x14:dataValidation type="list" allowBlank="1">'
    "<x14:formula1><xm:f>Values!$A$1:$A$2</xm:f></x14:formula1>"
    "<xm:sqref>C2:C4</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>"
)


def patch_part(path: Path, part: str, edit: Callable[[str], str]) -> None:
    """Rewrite one XML part of the workbook with ``edit(text) -> text``."""
    with zipfile.ZipFile(path) as src:
        parts = {name: src.read(name) for name in src.namelist()}
    parts[part] = edit(parts[part].decode("utf-8")).encode("utf-8")
    with zipfile.ZipFile(path, "w") as dst:
        for name, data in parts.items():
            dst.writestr(name, data)


def questions_workbook(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sec"
    sheet.append(["ID", "Item", "Choice", "Notes"])
    for n in range(1, 4):
        sheet.append([f"Q-{n}", f"Question number {n}?"])
    values = workbook.create_sheet("Values")
    values.append(["Yes"])
    values.append(["No"])
    workbook.save(path)
    patch_part(path, "xl/worksheets/sheet1.xml", lambda x: x.replace("</worksheet>", X14 + "</worksheet>"))
    return path


def test_x14_list_is_read_from_sheet_xml(tmp_path: Path) -> None:
    path = questions_workbook(tmp_path / "q.xlsx")
    assert read_x14_lists(path)["Sec"][0].sqref == "C2:C4"
    snap = load_snapshot(path)["Sec"]
    assert snap.list_values(3, range(2, 5)) == {"Yes", "No"}


def test_allowed_values_are_checked_against_x14_list(tmp_path: Path) -> None:
    snap = load_snapshot(questions_workbook(tmp_path / "q.xlsx"))["Sec"]

    def sheet_map(values: list[str]) -> SheetMap:
        return SheetMap.model_validate(
            {
                "sheet": "Sec",
                "header_row": 1,
                "question_col": "B",
                "answer_targets": {"short_answer_col": "C", "allowed_values": values, "text_col": "D"},
            }
        )

    assert check_sheet_map(sheet_map(["Yes", "No"]), snap).answer_targets.allowed_values == ["Yes", "No"]
    with pytest.raises(LayoutRejected, match="drop-down"):
        check_sheet_map(sheet_map(["Yes", "Maybe"]), snap)


def test_doctype_in_sheet_xml_is_refused(tmp_path: Path) -> None:
    path = questions_workbook(tmp_path / "q.xlsx")
    patch_part(path, "xl/worksheets/sheet1.xml", lambda x: '<!DOCTYPE a [<!ENTITY e "x">]>' + x)
    assert read_x14_lists(path) == {}


def marks_workbook(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Marks"
    sheet.append(["ID", "Question", "Yes", "No", "Not Applicable", "Notes"])
    for n in range(1, 4):
        sheet.append([f"M-{n}", f"Is control {n} in place?"])
    workbook.save(path)
    return path


def marks_map(**targets) -> SheetMap:
    base = {"text_col": "F", "choice_cols": {"Yes": "C", "No": "D", "Not Applicable": "E"}} | targets
    return SheetMap.model_validate(
        {"sheet": "Marks", "header_row": 1, "question_col": "B", "answer_targets": base}
    )


def test_mark_columns_are_accepted_when_titled_by_their_value(tmp_path: Path) -> None:
    snap = load_snapshot(marks_workbook(tmp_path / "m.xlsx"))["Marks"]
    checked = check_sheet_map(marks_map(), snap)
    assert checked.answer_targets.choice_cols == {"Yes": "C", "No": "D", "Not Applicable": "E"}


@pytest.mark.parametrize(
    ("targets", "message"),
    [
        ({"choice_cols": {"Yes": "C", "No": "C"}}, "also the"),
        ({"choice_cols": {"Yes": "C"}}, "at least two"),
        ({"choice_cols": {"Yes": "C", "Maybe": "D"}}, "not titled"),
        ({"short_answer_col": "C"}, "cannot both"),
        ({"choice_cols": {"Yes": "C", "No": "F"}}, "also the"),
    ],
)
def test_bad_mark_columns_are_rejected(tmp_path: Path, targets: dict, message: str) -> None:
    snap = load_snapshot(marks_workbook(tmp_path / "m.xlsx"))["Marks"]
    with pytest.raises(LayoutRejected, match=message):
        check_sheet_map(marks_map(**targets), snap)


class ReplyLLM:
    def __init__(self, reply: dict) -> None:
        self.reply = reply

    def complete_json(self, system: str, user: str) -> LLMReply:
        return LLMReply(content=json.dumps(self.reply), usage=LLMUsage(model="fake"))


def sheet_entry(name: str) -> dict:
    return {"sheet": name, "header_row": 1, "question_col": "B", "answer_targets": {"text_col": "C"}}


def test_sheet_repeating_another_sheet_is_not_asked_twice(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "Main"
    log = workbook.create_sheet("Log")
    for sheet in (workbook["Main"], log):
        sheet.append(["ID", "Question", "Answer"])
        for n in range(1, 5):
            sheet.append([f"Q-{n}", f"Does the vendor do thing {n}?"])
    workbook.save(tmp_path / "r.xlsx")
    llm = ReplyLLM({"sheets": [sheet_entry("Main"), sheet_entry("Log")]})
    result = resolve_xlsx(tmp_path / "r.xlsx", ManualLayout(), llm=llm, prompt="p")
    assert {q.location.sheet for q in result.questions} == {"Main"}
    assert len(result.questions) == 4
    log_map = next(m for m in result.layout.sheets if m.sheet == "Log")
    assert not log_map.has_questions
    assert log_map.repeat_of == "Main"


def test_repeat_detection_needs_a_large_overlap() -> None:
    def qs(*texts: str) -> list[Question]:
        return [
            Question(qid=f"q{i}", text=t, location=Location(sheet="s", row=i)) for i, t in enumerate(texts)
        ]

    found = find_repeat_sheets(
        {"A": qs("a", "b", "c", "d"), "B": qs("a", "b", "x", "y", "z"), "C": qs("a", "B ", "c")}
    )
    assert found == {"C": "A"}


def test_text_column_is_filled_when_the_model_gives_none(tmp_path: Path) -> None:
    path = questions_workbook(tmp_path / "q.xlsx")
    entry = {
        "sheet": "Sec",
        "header_row": 1,
        "question_col": "B",
        "answer_targets": {"short_answer_col": "C", "allowed_values": ["Yes", "No"], "text_col": None},
    }
    result = resolve_xlsx(path, ManualLayout(), llm=ReplyLLM({"sheets": [entry]}), prompt="p")
    sec = result.layout.sheets[0]
    assert sec.answer_targets.text_col == "E"


def test_failing_sheet_does_not_cost_the_others(tmp_path: Path) -> None:
    path = questions_workbook(tmp_path / "q.xlsx")
    bad = {"sheet": "Values", "header_row": 0, "question_col": "A", "answer_targets": {"text_col": "A"}}
    reply = {"sheets": [sheet_entry("Sec"), bad]}
    result = resolve_xlsx(path, ManualLayout(), llm=ReplyLLM(reply), prompt="p")
    assert result.layout.source == "llm"
    assert {q.location.sheet for q in result.questions} == {"Sec"}
    assert next(m for m in result.layout.sheets if m.sheet == "Values").left_out


def test_question_text_from_a_formula_is_read_and_export_keeps_the_formula(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Form"
    workbook.create_sheet("Hidden")["A1"] = "What is your uptime target?"
    sheet.append(["ID", "Question", "Answer"])
    sheet.append(["F-1", "=Hidden!A1"])
    path = tmp_path / "f.xlsx"
    workbook.save(path)

    def with_cached_value(xml: str) -> str:  # what Excel stores next to the formula
        xml = re.sub(r'(<c r="B2")[^>]*>', r'\1 t="str">', xml)
        return re.sub(r"(<f>[^<]*</f>)<v\s*/>|(<f>[^<]*</f>)<v></v>", r"\1\2<v>cached text</v>", xml)

    patch_part(path, "xl/worksheets/sheet1.xml", with_cached_value)
    entry = sheet_entry("Form")
    result = resolve_xlsx(path, ManualLayout(), llm=ReplyLLM({"sheets": [entry]}), prompt="p")
    question = result.questions[0]
    assert question.text == "cached text"
    answers = {question.qid: Answer(qid=question.qid, text="99.9%", status="answered", confidence=0.9)}
    out = tmp_path / "out.xlsx"
    export_xlsx(path, result.questions, answers, out)
    verify_xlsx(path, out, result.questions, answers)


def test_text_column_that_is_a_mark_column_is_moved_to_a_free_one(tmp_path: Path) -> None:
    path = marks_workbook(tmp_path / "m.xlsx")
    entry = {
        "sheet": "Marks",
        "header_row": 1,
        "question_col": "B",
        "answer_targets": {"text_col": "C", "choice_cols": {"Yes": "C", "No": "D", "Not Applicable": "E"}},
    }
    result = resolve_xlsx(path, ManualLayout(), llm=ReplyLLM({"sheets": [entry]}), prompt="p")
    targets = result.layout.sheets[0].answer_targets
    assert targets.choice_cols == {"Yes": "C", "No": "D", "Not Applicable": "E"}
    assert targets.text_col == "G"


def test_text_column_with_a_drop_down_list_becomes_the_short_answer(tmp_path: Path) -> None:
    path = questions_workbook(tmp_path / "q.xlsx")
    entry = {
        "sheet": "Sec",
        "header_row": 1,
        "question_col": "B",
        "answer_targets": {"text_col": "C"},
    }
    result = resolve_xlsx(path, ManualLayout(), llm=ReplyLLM({"sheets": [entry]}), prompt="p")
    targets = result.layout.sheets[0].answer_targets
    assert (targets.short_answer_col, targets.text_col) == ("C", "E")
