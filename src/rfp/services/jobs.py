"""Questionnaire jobs: upload, background answering, review of answers and export.

A job is one uploaded XLSX/DOCX/PDF. Its questions are extracted when the job is created (so a bad
file is rejected at once); a worker thread then answers them one by one, committing each answer
as it is produced so that progress (``done`` of ``total``) can be polled. A failure on one question
turns that question into ``needs_review`` with a note and does not stop the job. Jobs left
``queued`` or ``running`` by a restart are resumed: only still ``pending`` questions are processed.

Each item keeps the model's draft (``model_text``) and the current ``text``/``status``, which a
reviewer may change; export always writes the current values.
"""

import logging
import re
import tempfile
import time
from collections.abc import Callable
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from typing import Literal

import psycopg
from openpyxl import load_workbook
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from pydantic import BaseModel, Field

from rfp.answering import GeneratedAnswer, answer_question
from rfp.audit import log_action, log_model_request
from rfp.config import Settings
from rfp.embeddings import Embedder
from rfp.export import export_answers
from rfp.llm import ChatClient
from rfp.models import Answer, AnswerStatus, Citation, Location, Question, RetrievedChunk
from rfp.questions import extract_docx_questions, extract_pdf_questions, extract_xlsx_questions
from rfp.retrieval import retrieve
from rfp.services.errors import ConflictError, InvalidInputError, NotFoundError

log = logging.getLogger(__name__)

JobStatus = Literal["queued", "running", "done", "failed"]
ItemStatus = Literal["pending", "answered", "approved", "edited", "needs_review", "no_data"]
ReviewStatus = Literal["approved", "edited", "needs_review", "no_data"]

JOB_SUFFIXES = (".xlsx", ".docx", ".pdf")
MEDIA_TYPES = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
FAILED_ANSWER_TEXT = "Automatic answering failed; this question needs a manual answer."
MAX_NOTE_CHARS = 500
_COLUMN_RE = re.compile(r"^[A-Za-z]{1,3}$")


class JobParams(BaseModel):
    """Where the questions are in an XLSX file; ignored for DOCX and PDF."""

    sheet: str | None = Field(default=None, description="Sheet name; the first sheet when omitted")
    question_col: str = Field(default="B", description="Column letter holding the questions")
    answer_col: str = Field(default="C", description="Column letter the answers are written to")
    header_row: int = Field(default=1, ge=0, description="Questions are read from the rows below this one")
    number_col: str | None = Field(default=None, description="Optional column letter with item numbers")


class JobSummary(BaseModel):
    id: int
    filename: str
    status: JobStatus
    total: int = Field(description="Number of questions found in the file")
    done: int = Field(description="Number of questions processed so far")
    error: str | None = None
    created_by: str
    created_at: datetime
    started_at: datetime | None = None
    finished_at: datetime | None = None


class JobItem(BaseModel):
    qid: str
    number: str | None = Field(default=None, description="Item number in the source document, if any")
    question: str
    location: Location
    status: ItemStatus
    confidence: float
    text: str = Field(description="Current answer text (the model's draft until a reviewer edits it)")
    model_text: str | None = Field(default=None, description="The model's original draft")
    citations: list[Citation]
    note: str | None = None
    reviewed_by: str | None = None
    reviewed_at: datetime | None = None


class JobDetail(JobSummary):
    items: list[JobItem]


class ExportFile(BaseModel):
    filename: str
    content: bytes
    media_type: str


# --- creation -----------------------------------------------------------------------------------


def create_job(
    conn: psycopg.Connection, username: str, filename: str, data: bytes, params: JobParams
) -> JobSummary:
    """Extract the questions of an uploaded file and store the job as ``queued``."""
    name = Path(filename).name
    suffix = Path(name).suffix.lower()
    if suffix not in JOB_SUFFIXES:
        raise InvalidInputError(f"Unsupported file type '{suffix}'; use XLSX, DOCX or PDF.")
    if not data:
        raise InvalidInputError("The file is empty.")
    questions = _extract_questions(name, data, params)
    if not questions:
        raise InvalidInputError("No questions were found in the file.")
    with conn.transaction():
        job_id = conn.execute(
            """
            INSERT INTO jobs (filename, source_file, params, total, created_by)
            VALUES (%s, %s, %s, %s, %s) RETURNING id
            """,
            (name, data, Jsonb(params.model_dump()), len(questions), username),
        ).fetchone()[0]
        with conn.cursor() as cur:
            cur.executemany(
                "INSERT INTO job_items (job_id, qid, ord, question) VALUES (%s, %s, %s, %s)",
                [(job_id, q.qid, order, Jsonb(q.model_dump())) for order, q in enumerate(questions)],
            )
        log_action(
            conn, username, "job.create", "job", job_id, {"filename": name, "questions": len(questions)}
        )
    return get_job_summary(conn, job_id)


def _extract_questions(name: str, data: bytes, params: JobParams) -> list[Question]:
    suffix = Path(name).suffix.lower()
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / f"input{suffix}"
        path.write_bytes(data)
        try:
            if suffix == ".xlsx":
                return _extract_xlsx(path, params)
            if suffix == ".docx":
                return extract_docx_questions(path)
            return extract_pdf_questions(path)
        except InvalidInputError:
            raise
        except Exception as exc:
            log.warning("Cannot read questions from %s: %s", name, exc)
            raise InvalidInputError(f"Cannot read questions from '{name}': {exc}") from exc


def _extract_xlsx(path: Path, params: JobParams) -> list[Question]:
    for column in (params.question_col, params.answer_col, params.number_col or "A"):
        if not _COLUMN_RE.match(column):
            raise InvalidInputError(f"'{column}' is not a column letter.")
    sheet = params.sheet
    workbook = load_workbook(path, read_only=True)
    try:
        if sheet is None:
            sheet = workbook.sheetnames[0]
        elif sheet not in workbook.sheetnames:
            raise InvalidInputError(
                f"Sheet '{sheet}' not found; available: {', '.join(workbook.sheetnames)}."
            )
    finally:
        workbook.close()
    return extract_xlsx_questions(
        path, sheet, params.question_col, params.answer_col, params.header_row, params.number_col
    )


# --- reading ------------------------------------------------------------------------------------

_SUMMARY_SQL = """
    SELECT j.id, j.filename, j.status, j.total, j.error, j.created_by, j.created_at, j.started_at,
           j.finished_at,
           (SELECT count(*) FROM job_items i WHERE i.job_id = j.id AND i.status <> 'pending') AS done
    FROM jobs j
"""


def get_job_summary(conn: psycopg.Connection, job_id: int) -> JobSummary:
    with conn.cursor(row_factory=dict_row) as cur:
        row = cur.execute(f"{_SUMMARY_SQL} WHERE j.id = %s", (job_id,)).fetchone()
    if row is None:
        raise NotFoundError(f"Job {job_id} not found.")
    return JobSummary.model_validate(row)


def list_jobs(conn: psycopg.Connection, limit: int, offset: int) -> tuple[list[JobSummary], int]:
    """Newest first; returns the page and the total number of jobs."""
    with conn.cursor(row_factory=dict_row) as cur:
        rows = cur.execute(
            f"{_SUMMARY_SQL} ORDER BY j.id DESC LIMIT %s OFFSET %s", (limit, offset)
        ).fetchall()
        total = cur.execute("SELECT count(*) AS n FROM jobs").fetchone()
    return [JobSummary.model_validate(row) for row in rows], total["n"] if total else 0


def get_job(conn: psycopg.Connection, job_id: int) -> JobDetail:
    summary = get_job_summary(conn, job_id)
    return JobDetail(**summary.model_dump(), items=_load_items(conn, job_id))


def _load_items(conn: psycopg.Connection, job_id: int, qid: str | None = None) -> list[JobItem]:
    where, params = ("AND qid = %s", [job_id, qid]) if qid else ("", [job_id])
    with conn.cursor(row_factory=dict_row) as cur:
        rows = cur.execute(
            f"""
            SELECT qid, question, status, confidence, text, model_text, citations, note,
                   reviewed_by, reviewed_at
            FROM job_items WHERE job_id = %s {where} ORDER BY ord
            """,
            params,
        ).fetchall()
    items = []
    for row in rows:
        question = Question.model_validate(row.pop("question"))
        items.append(
            JobItem(
                **row,
                number=question.number,
                question=question.text,
                location=question.location,
            )
        )
    return items


# --- processing ---------------------------------------------------------------------------------


def process_job(
    conn: psycopg.Connection, embedder: Embedder, llm: ChatClient, settings: Settings, job_id: int
) -> None:
    """Answer every still pending question of the job; mark the job ``done`` (or ``failed`` on a crash)."""
    try:
        conn.execute(
            """
            UPDATE jobs SET status = 'running', started_at = COALESCE(started_at, now()), error = NULL
            WHERE id = %s
            """,
            (job_id,),
        )
        rows = conn.execute(
            "SELECT question FROM job_items WHERE job_id = %s AND status = 'pending' ORDER BY ord", (job_id,)
        ).fetchall()
        for (raw,) in rows:
            _process_item(conn, embedder, llm, settings, job_id, Question.model_validate(raw))
        conn.execute("UPDATE jobs SET status = 'done', finished_at = now() WHERE id = %s", (job_id,))
    except Exception as exc:
        log.exception("Job %s failed", job_id)
        conn.execute(
            "UPDATE jobs SET status = 'failed', error = %s, finished_at = now() WHERE id = %s",
            (_short(f"{type(exc).__name__}: {exc}"), job_id),
        )


def _process_item(
    conn: psycopg.Connection,
    embedder: Embedder,
    llm: ChatClient,
    settings: Settings,
    job_id: int,
    question: Question,
) -> None:
    started = time.perf_counter()
    chunks: list[RetrievedChunk] = []
    generated: GeneratedAnswer | None = None
    error: str | None = None
    try:
        chunks = retrieve(conn, embedder, question.text)
        generated = answer_question(question, chunks, llm, settings)
    except Exception as exc:
        error = _short(f"{type(exc).__name__}: {exc}")
        log.warning("Job %s, question %s failed: %s", job_id, question.qid, error)
    duration = time.perf_counter() - started
    log_model_request(
        conn,
        job_id=job_id,
        qid=question.qid,
        question=question.text,
        chunks=chunks,
        generated=generated,
        duration_s=duration,
        error=error,
    )
    if generated is None:
        text, status, confidence, citations, note = FAILED_ANSWER_TEXT, "needs_review", 0.0, [], error
    else:
        answer = generated.answer
        text, status, confidence, citations = answer.text, answer.status, answer.confidence, answer.citations
        note = (
            f"{generated.rejected_citations} model citation(s) rejected: not found in the fragments."
            if generated.rejected_citations
            else None
        )
    conn.execute(
        """
        UPDATE job_items SET status = %s, text = %s, confidence = %s, citations = %s,
                             model_text = %s, model_status = %s, note = %s
        WHERE job_id = %s AND qid = %s
        """,
        (
            status,
            text,
            confidence,
            Jsonb([c.model_dump() for c in citations]),
            text if generated else None,
            status if generated else None,
            note,
            job_id,
            question.qid,
        ),
    )


def _short(message: str) -> str:
    return message if len(message) <= MAX_NOTE_CHARS else message[: MAX_NOTE_CHARS - 1] + "…"


class JobRunner:
    """Runs jobs on a small thread pool inside the process; each job uses its own connection."""

    def __init__(
        self,
        connect: Callable[[], psycopg.Connection],
        embedder: Embedder,
        llm: ChatClient,
        settings: Settings,
    ) -> None:
        self._connect = connect
        self._embedder = embedder
        self._llm = llm
        self._settings = settings
        self._executor = ThreadPoolExecutor(max_workers=settings.job_workers, thread_name_prefix="job")

    def submit(self, job_id: int) -> None:
        self._executor.submit(self._run, job_id)

    def resume_unfinished(self, conn: psycopg.Connection) -> int:
        """Re-queue jobs that a restart interrupted; returns how many. Assumes a single app process."""
        ids = [
            row[0]
            for row in conn.execute("SELECT id FROM jobs WHERE status IN ('queued', 'running') ORDER BY id")
        ]
        for job_id in ids:
            self.submit(job_id)
        return len(ids)

    def shutdown(self) -> None:
        self._executor.shutdown(wait=False, cancel_futures=True)

    def _run(self, job_id: int) -> None:
        try:
            conn = self._connect()
        except Exception:
            log.exception("Job %s cannot start: no database connection", job_id)
            return
        try:
            process_job(conn, self._embedder, self._llm, self._settings, job_id)
        finally:
            conn.close()


# --- review and export --------------------------------------------------------------------------


def review_item(
    conn: psycopg.Connection,
    job_id: int,
    qid: str,
    username: str,
    text: str | None = None,
    status: ReviewStatus | None = None,
    note: str | None = None,
) -> JobItem:
    """Change an answer's text and/or status and record who did it. A new text alone means ``edited``."""
    if text is None and status is None and note is None:
        raise InvalidInputError("Nothing to change: give text, status or note.")
    if text is not None and not text.strip():
        raise InvalidInputError("The answer text must not be empty.")
    current = _load_items(conn, job_id, qid)
    if not current:
        raise NotFoundError(f"Question {qid} not found in job {job_id}.")
    if current[0].status == "pending":
        raise ConflictError("The question has not been processed yet.")
    if status is None and text is not None:
        status = "edited"
    with conn.transaction():
        conn.execute(
            """
            UPDATE job_items SET text = COALESCE(%s, text), status = COALESCE(%s, status),
                                 note = COALESCE(%s, note), reviewed_by = %s, reviewed_at = now()
            WHERE job_id = %s AND qid = %s
            """,
            (text.strip() if text is not None else None, status, note, username, job_id, qid),
        )
        log_action(
            conn,
            username,
            "answer.edit",
            "job_item",
            f"{job_id}/{qid}",
            {"status": status, "text_changed": text is not None},
        )
    return _load_items(conn, job_id, qid)[0]


def export_job(conn: psycopg.Connection, job_id: int, username: str) -> ExportFile:
    """The source file with the current answers written into it (a PDF input gives a DOCX report)."""
    row = conn.execute("SELECT filename, source_file, status FROM jobs WHERE id = %s", (job_id,)).fetchone()
    if row is None:
        raise NotFoundError(f"Job {job_id} not found.")
    filename, source_file, job_status = row
    if job_status in ("queued", "running"):
        raise ConflictError("The job is still being processed.")
    items = [item for item in _load_items(conn, job_id) if item.status != "pending"]
    questions = [Question(qid=i.qid, text=i.question, number=i.number, location=i.location) for i in items]
    answers = {
        i.qid: Answer(
            qid=i.qid,
            text=i.text,
            status=_export_status(i.status),
            confidence=min(1.0, max(0.0, i.confidence)),
            citations=i.citations,
        )
        for i in items
    }
    suffix = Path(filename).suffix.lower()
    out_suffix = ".docx" if suffix == ".pdf" else suffix
    with tempfile.TemporaryDirectory() as tmp:
        source = Path(tmp) / f"source{suffix}"
        dest = Path(tmp) / f"answers{out_suffix}"
        source.write_bytes(source_file)
        export_answers(source, questions, answers, dest)
        content = dest.read_bytes()
    log_action(conn, username, "job.export", "job", job_id, {"filename": filename})
    return ExportFile(
        filename=f"{Path(filename).stem}-answers{out_suffix}",
        content=content,
        media_type=MEDIA_TYPES[out_suffix],
    )


def _export_status(status: ItemStatus) -> AnswerStatus:
    """Collapse review states into the export's vocabulary: reviewed and drafted answers are 'answered'."""
    if status == "no_data":
        return "no_data"
    if status == "needs_review":
        return "needs_review"
    return "answered"
