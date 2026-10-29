"""A read-only picture of an XLSX workbook: cell texts, merged ranges and drop-down lists.

It feeds the layout call (as compact text with cell addresses) and the checks of the model's map.
"""

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from rfp.questions.xlsx_x14 import read_x14_lists

HEAD_ROWS = 25
TAIL_ROWS = 5
MIDDLE_ROWS = 5
MAX_CELL_CHARS = 100
MAX_COLUMNS = 40
MAX_LIST_VALUES = 40

_REFERENCE_RE = re.compile(
    r"^(?:'?(?P<sheet>[^'!]+)'?!)?(?P<range>\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)$"
)


@dataclass(frozen=True)
class Area:
    """A rectangle of cells, 1-based and inclusive."""

    min_row: int
    min_col: int
    max_row: int
    max_col: int

    def contains(self, row: int, col: int) -> bool:
        return self.min_row <= row <= self.max_row and self.min_col <= col <= self.max_col

    @property
    def label(self) -> str:
        start = f"{get_column_letter(self.min_col)}{self.min_row}"
        end = f"{get_column_letter(self.max_col)}{self.max_row}"
        return start if start == end else f"{start}:{end}"


@dataclass(frozen=True)
class ListRule:
    """A data-validation drop-down list over ``areas``; ``values`` is None when it cannot be resolved."""

    areas: tuple[Area, ...]
    values: tuple[str, ...] | None

    def covers_column(self, col: int, rows: range) -> bool:
        return any(
            a.min_col <= col <= a.max_col and a.min_row <= rows.stop - 1 and a.max_row >= rows.start
            for a in self.areas
        )


@dataclass
class SheetSnapshot:
    name: str
    max_row: int
    max_col: int
    hidden: bool = False
    cells: dict[tuple[int, int], str] = field(default_factory=dict)
    merged: list[Area] = field(default_factory=list)
    lists: list[ListRule] = field(default_factory=list)

    def text(self, row: int, col: int) -> str:
        return self.cells.get((row, col), "")

    def merged_at(self, row: int, col: int) -> Area | None:
        return next((a for a in self.merged if a.contains(row, col)), None)

    def list_values(self, col: int, rows: range) -> set[str] | None:
        """Union of the values of the lists that cover column ``col`` in ``rows``; None without a list."""
        found = [rule for rule in self.lists if rule.covers_column(col, rows)]
        if not found:
            return None
        values: set[str] = set()
        for rule in found:
            values.update(rule.values or ())
        return values


def load_snapshot(path: Path) -> dict[str, SheetSnapshot]:
    """All sheets of the workbook, in order. Cached formula results stand in for formulas.

    Drop-down lists kept in the x14 extension, which openpyxl discards, are read from the sheet XML.
    """
    workbook = load_workbook(path, data_only=True)
    try:
        snapshots = {ws.title: _sheet_snapshot(workbook, ws) for ws in workbook.worksheets}
        for name, extension_lists in read_x14_lists(path).items():
            if name not in snapshots:
                continue
            for item in extension_lists:
                areas = tuple(Area(*_bounds(part)) for part in item.sqref.split())
                snapshots[name].lists.append(
                    ListRule(areas, _list_values(workbook, workbook[name], item.formula))
                )
        return snapshots
    finally:
        workbook.close()


def _bounds(reference: str) -> tuple[int, int, int, int]:
    """(min_row, min_col, max_row, max_col) of ``C5`` or ``C5:D9``."""
    min_col, min_row, max_col, max_row = range_boundaries(reference.replace("$", ""))
    return min_row, min_col, max_row, max_col


def _sheet_snapshot(workbook, ws) -> SheetSnapshot:
    snapshot = SheetSnapshot(ws.title, ws.max_row or 0, ws.max_column or 0, ws.sheet_state != "visible")
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip():
                snapshot.cells[(cell.row, cell.column)] = str(cell.value).strip()
    for merged in ws.merged_cells.ranges:
        snapshot.merged.append(Area(merged.min_row, merged.min_col, merged.max_row, merged.max_col))
    for rule in ws.data_validations.dataValidation:
        if rule.type != "list":
            continue
        areas = tuple(Area(r.min_row, r.min_col, r.max_row, r.max_col) for r in rule.sqref.ranges)
        snapshot.lists.append(ListRule(areas, _list_values(workbook, ws, rule.formula1)))
    return snapshot


def _list_values(workbook, ws, formula: str | None) -> tuple[str, ...] | None:
    if not formula:
        return None
    formula = formula.strip()
    if formula.startswith('"') and formula.endswith('"') and len(formula) >= 2:
        return tuple(v.strip() for v in formula[1:-1].split(",") if v.strip())
    match = _REFERENCE_RE.match(formula)
    if match is None:
        return None
    target = workbook[match["sheet"]] if match["sheet"] in workbook.sheetnames else None
    if match["sheet"] is None:
        target = ws
    if target is None:
        return None
    min_col, min_row, max_col, max_row = range_boundaries(match["range"].replace("$", ""))
    values = []
    for row in target.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if cell.value is not None and str(cell.value).strip():
                values.append(str(cell.value).strip())
    return tuple(values)


def render_sheet(sheet: SheetSnapshot) -> str:
    """Compact text for the model: addresses and values of the first rows and a few later ones."""
    lines = [f'SHEET "{sheet.name}": {sheet.max_row} rows, {sheet.max_col} columns']
    filled = sorted({row for row, _ in sheet.cells})
    if not filled:
        return "\n".join([*lines, "(empty sheet)"])
    if sheet.merged:
        shown = ", ".join(a.label for a in sheet.merged[:20])
        lines.append(
            f"MERGED CELLS: {shown}"
            + (f" (+{len(sheet.merged) - 20} more)" if len(sheet.merged) > 20 else "")
        )
    for rule in sheet.lists:
        where = ", ".join(a.label for a in rule.areas[:6])
        values = (
            ", ".join(rule.values[:MAX_LIST_VALUES]) if rule.values is not None else "(taken from elsewhere)"
        )
        lines.append(f"DROP-DOWN LIST at {where}: {values}")
    lines.append("ROWS (empty cells omitted):")
    head = [r for r in filled if r <= HEAD_ROWS]
    rest = [r for r in filled if r > HEAD_ROWS]
    middle = (
        rest[len(rest) // 2 - MIDDLE_ROWS // 2 :][:MIDDLE_ROWS] if len(rest) > TAIL_ROWS + MIDDLE_ROWS else []
    )
    tail = rest[-TAIL_ROWS:] if rest else []
    previous = 0
    for row in dict.fromkeys([*head, *middle, *tail]):
        if previous and row - previous > 1:
            lines.append(f"... rows {previous + 1}-{row - 1} not shown ...")
        lines.append(_render_row(sheet, row))
        previous = row
    if previous < sheet.max_row:
        lines.append(f"... rows {previous + 1}-{sheet.max_row} not shown ...")
    return "\n".join(lines)


def _render_row(sheet: SheetSnapshot, row: int) -> str:
    parts = []
    for col in range(1, min(sheet.max_col, MAX_COLUMNS) + 1):
        value = sheet.text(row, col)
        if value:
            value = " ".join(value.split())
            if len(value) > MAX_CELL_CHARS:
                value = value[: MAX_CELL_CHARS - 1] + "…"
            parts.append(f'{get_column_letter(col)}{row}="{value}"')
    return " | ".join(parts)
