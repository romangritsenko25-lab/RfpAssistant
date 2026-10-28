"""Where the questions and answers live in an XLSX questionnaire.

The layout is found from the header row (Russian, Kazakh and English names); any part can be
given explicitly and then wins over the detected one.
"""

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from rfp.services.errors import InvalidInputError

HEADER_SCAN_ROWS = 10
COLUMN_RE = re.compile(r"^[A-Za-z]{1,3}$")

_QUESTION = {
    "question",
    "questions",
    "requirement",
    "requirements",
    "вопрос",
    "вопросы",
    "требование",
    "требования",
}
_QUESTION_KK = {"сұрақ", "сұрақтар", "талап", "талаптар"}
_ANSWER = {"answer", "answers", "response", "vendorresponse", "vendoranswer", "supplierresponse"}
_ANSWER_RU = {"ответ", "ответы", "ответпоставщика", "ответисполнителя"}
_ANSWER_KK = {"жауап", "жауаптар", "жеткізушіжауабы"}
_NUMBER = {"id", "no", "number", "#", "№", "номер", "п/п", "№п/п", "нөмір", "reqid"}


@dataclass(frozen=True)
class XlsxLayout:
    """Resolved columns of one sheet. ``answer_header`` is set only when the answer column did not
    exist and must get this title in the header row when the answers are exported."""

    sheet: str
    header_row: int
    question_col: str
    answer_col: str
    number_col: str | None = None
    answer_header: str | None = None


@dataclass(frozen=True)
class _Header:
    row: int
    question: int
    answer: int | None
    number: int | None


def detect_xlsx_layout(
    path: Path,
    sheet: str | None = None,
    *,
    header_row: int | None = None,
    question_col: str | None = None,
    answer_col: str | None = None,
    number_col: str | None = None,
) -> XlsxLayout:
    """Find the header row and the question, answer and number columns of ``sheet``.

    The header row is the first of the top 10 rows that has a question-like title. Values passed in
    are kept as they are; only the missing ones are detected. Without an answer title the answer
    goes to the first empty column right of the header. Raises ``InvalidInputError`` when the
    question column cannot be determined or the choice is unsafe (answers over questions).
    """
    question_col, answer_col, number_col = (_column(c) for c in (question_col, answer_col, number_col))
    needs_scan = question_col is None or answer_col is None or header_row is None
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is None:
            sheet = workbook.sheetnames[0]
        elif sheet not in workbook.sheetnames:
            raise InvalidInputError(
                f"Sheet '{sheet}' not found; available: {', '.join(workbook.sheetnames)}."
            )
        rows: list[tuple] = []
        if needs_scan or number_col is None:
            last_row = max(HEADER_SCAN_ROWS, header_row or 0)
            rows = list(workbook[sheet].iter_rows(min_row=1, max_row=last_row, values_only=True))
    finally:
        workbook.close()

    found = _find_header(rows, header_row)
    if question_col is None:
        if found is None:
            raise InvalidInputError(
                f"Cannot find the question column: no header such as 'Question', 'Вопрос' or 'Сұрақ' "
                f"in the first {HEADER_SCAN_ROWS} rows. Specify the columns manually."
            )
        question_col = get_column_letter(found.question)
    if header_row is None:
        header_row = found.row if found else 1
    if number_col is None and found and found.number:
        number_col = get_column_letter(found.number)

    answer_header = None
    if answer_col is None:
        if found and found.answer:
            answer_col = get_column_letter(found.answer)
        else:
            header_cells = rows[header_row - 1] if 1 <= header_row <= len(rows) else ()
            last_filled = max((i for i, v in enumerate(header_cells, 1) if _text(v)), default=0)
            answer_col = get_column_letter(max(last_filled, column_index_from_string(question_col)) + 1)
            if last_filled:
                answer_header = _answer_title(header_cells, column_index_from_string(question_col))

    if answer_col == question_col:
        raise InvalidInputError("The answer column must differ from the question column.")
    if number_col == answer_col:
        raise InvalidInputError("The number column must differ from the answer column.")
    return XlsxLayout(sheet, header_row, question_col, answer_col, number_col, answer_header)


def _column(value: str | None) -> str | None:
    value = (value or "").strip()
    if not value:
        return None
    if not COLUMN_RE.match(value):
        raise InvalidInputError(f"'{value}' is not a column letter.")
    return value.upper()


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _normal(value: object) -> str:
    """Lower case, no whitespace, no trailing ':'/'*' — so 'Vendor Response:' equals 'vendorresponse'."""
    return "".join(_text(value).casefold().replace("ё", "е").split()).rstrip(":*.")  # noqa: RUF001


def _find_header(rows: list[tuple], only_row: int | None) -> _Header | None:
    for number, cells in enumerate(rows, 1):
        if only_row is not None and number != only_row:
            continue
        if only_row is None and number > HEADER_SCAN_ROWS:
            break
        roles: dict[str, int] = {}
        for index, value in enumerate(cells, 1):
            name = _normal(value)
            for role, names in (
                ("question", _QUESTION | _QUESTION_KK),
                ("answer", _ANSWER | _ANSWER_RU | _ANSWER_KK),
                ("number", _NUMBER),
            ):
                if name in names:
                    roles.setdefault(role, index)
        if "question" in roles:
            return _Header(number, roles["question"], roles.get("answer"), roles.get("number"))
    return None


def _answer_title(header_cells: tuple, question_index: int) -> str:
    """Title for a new answer column, in the language of the question header."""
    name = _normal(header_cells[question_index - 1]) if question_index <= len(header_cells) else ""
    if name in _QUESTION_KK:
        return "Жауап"
    if name in _QUESTION and re.search("[а-я]", name):  # noqa: RUF001
        return "Ответ"
    return "Answer"
