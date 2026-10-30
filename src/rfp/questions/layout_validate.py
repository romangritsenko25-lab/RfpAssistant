"""Hard checks of a sheet map against the workbook itself.

The model proposes; this code decides. A map that breaks any rule is rejected as a whole, and the
caller falls back to header rules.
"""

from openpyxl.utils import column_index_from_string, get_column_letter

from rfp.questions.layout_map import SheetMap
from rfp.questions.xlsx_layout import answer_title
from rfp.questions.xlsx_snapshot import SheetSnapshot

MAX_PROBLEMS = 5
MAX_TITLE_CHARS = 60
# Titles of a column made for a free-text reply; the existing one is preferred to a new column.
_RESPONSE_TITLES = (
    "notes",
    "note",
    "comment",
    "response",
    "explanation",
    "answer",
    "remark",
    "ответ",
    "комментар",
    "примечан",
    "пояснен",
    "жауап",
    "түсініктеме",
    "ескертпе",
)


class LayoutRejected(ValueError):
    """The map is unsafe or does not match the file; the message lists what is wrong."""


def check_sheet_map(sheet_map: SheetMap, snap: SheetSnapshot) -> SheetMap:
    """Return the map completed from the file (rows, skipped sections, new column title) or raise.

    Rules: the answer columns differ from the question and number columns and from each other; the
    declared rows lie inside the sheet and the question cells hold text; there is at least one
    question; answer cells never hold or share a merged cell with their question; an answer column
    that is already mostly filled is refused; allowed values come from the sheet's drop-down list;
    mark columns (``choice_cols``) are at least two, distinct, empty, not combined with a short answer
    column, and each is titled with its value above the first question.
    """
    targets = sheet_map.answer_targets
    assert targets is not None and sheet_map.question_col is not None  # guaranteed by SheetMap
    where = f"sheet '{snap.name}'"
    problems: list[str] = []

    named = [("question", sheet_map.question_col), ("number", sheet_map.number_col)]
    named += [("short answer", targets.short_answer_col), ("answer text", targets.text_col)]
    named += [("source", targets.source_col)]
    named += [(f"mark '{value}'", col) for value, col in (targets.choice_cols or {}).items()]
    seen: dict[str, str] = {}
    for role, col in named:
        if col is None:
            continue
        if col in seen:
            problems.append(f"{where}: the {role} column {col} is also the {seen[col]} column")
        seen.setdefault(col, role)

    question_idx = column_index_from_string(sheet_map.question_col)
    header = sheet_map.header_row or 0
    first = sheet_map.first_question_row or header + 1
    last = sheet_map.last_question_row or snap.max_row
    if header > snap.max_row:
        problems.append(f"{where}: header row {header} is below the last row {snap.max_row}")
    if last > snap.max_row:
        problems.append(f"{where}: last question row {last} is below the last row {snap.max_row}")
    if first <= header:
        problems.append(f"{where}: the first question row {first} is not below the header row {header}")
    if first > last:
        problems.append(f"{where}: the first question row {first} is after the last one {last}")
    outside = [r for r in sheet_map.section_header_rows if not 1 <= r <= snap.max_row]
    if outside:
        problems.append(f"{where}: section header rows {outside} are outside the sheet")
    problems += _check_choice_cols(snap, targets, header, sheet_map.first_question_row or header + 1)
    if problems:
        raise LayoutRejected("; ".join(problems[:MAX_PROBLEMS]))

    section = {r for r in sheet_map.section_header_rows if first <= r <= last}
    for row in range(first, last + 1):
        merged = snap.merged_at(row, question_idx)
        if snap.text(row, question_idx) and merged is not None and merged.max_col > merged.min_col:
            section.add(row)  # a title stretched across the row is never a question
    question_rows = [r for r in range(first, last + 1) if r not in section and snap.text(r, question_idx)]
    if not question_rows:
        problems.append(f"{where}: no question text in column {sheet_map.question_col}, rows {first}-{last}")
    for edge, label in ((sheet_map.first_question_row, "first"), (sheet_map.last_question_row, "last")):
        if edge is not None and edge not in section and not snap.text(edge, question_idx):
            problems.append(f"{where}: the {label} question row {edge} has no question text")
    if problems:
        raise LayoutRejected("; ".join(problems[:MAX_PROBLEMS]))

    targets = _existing_text_column(snap, sheet_map, targets, header, question_rows)
    for col in targets.columns():
        # A drop-down column may already hold choices (a filled copy); nothing is lost by listing it.
        listed = col == targets.short_answer_col and snap.list_values(
            column_index_from_string(col), range(first, last + 1)
        )
        problems += _check_target_column(snap, col, question_idx, question_rows, bool(listed))
    allowed, allowed_ok = _check_allowed_values(
        snap, targets.short_answer_col, targets.allowed_values, first, last
    )
    if not allowed_ok:
        problems.append(
            f"{where}: allowed values of column {targets.short_answer_col} are not in its drop-down list"
        )
    if problems:
        raise LayoutRejected("; ".join(problems[:MAX_PROBLEMS]))

    header_cells = tuple(snap.text(header, c) or None for c in range(1, snap.max_col + 1))
    text_idx = column_index_from_string(targets.text_col)
    new_title = (
        answer_title(header_cells, question_idx) if header >= 1 and not snap.text(header, text_idx) else None
    )
    new_targets = targets.model_copy(update={"allowed_values": allowed or None})
    return sheet_map.model_copy(
        update={
            "header_row": header,
            "first_question_row": question_rows[0],
            "last_question_row": question_rows[-1],
            "section_header_rows": sorted(section),
            "skipped_rows": sum(1 for r in section if snap.text(r, question_idx)),
            "answer_targets": new_targets,
            "answer_header": new_title,
        }
    )


def _is_response_title(title: str) -> bool:
    folded = title.casefold()
    return 0 < len(folded) <= MAX_TITLE_CHARS and any(word in folded for word in _RESPONSE_TITLES)


def _existing_text_column(
    snap: SheetSnapshot, sheet_map: SheetMap, targets, header: int, question_rows: list[int]
):
    """Targets whose text column is the sheet's own empty response column, when there is one.

    The model tends to name a new column to the right although the table has a column titled
    "Comments" or "Response" that is empty in the question rows. A drop-down column, and any column
    already used for another part of the answer, is not a candidate.
    """
    if header < 1:
        return targets
    taken = {sheet_map.question_col, sheet_map.number_col, targets.short_answer_col, targets.source_col}
    taken |= set((targets.choice_cols or {}).values())
    current = snap.text(header, column_index_from_string(targets.text_col))
    if _is_response_title(current):
        return targets
    for index in range(1, snap.max_col + 1):
        col = get_column_letter(index)
        if col in taken or not _is_response_title(snap.text(header, index)):
            continue
        if snap.list_values(index, range(header + 1, snap.max_row + 1)):
            continue
        if any(snap.text(row, index) for row in question_rows):
            continue
        return targets.model_copy(update={"text_col": col})
    return targets


def _normal(text: str) -> str:
    return "".join(ch for ch in text.casefold() if ch.isalnum())


def _check_choice_cols(snap: SheetSnapshot, targets, header: int, first: int) -> list[str]:
    """Rules for an answer given as a mark in one of several columns; empty when there are none."""
    choices = targets.choice_cols
    if not choices:
        return []
    where = f"sheet '{snap.name}'"
    if len(choices) < 2:
        return [f"{where}: a mark answer needs at least two columns"]
    if targets.short_answer_col:
        return [f"{where}: the short answer column and mark columns cannot both be used"]
    problems = []
    for value, col in choices.items():
        index = column_index_from_string(col)
        titles = [_normal(snap.text(row, index)) for row in range(1, max(header, first - 1) + 1)]
        wanted = _normal(value)
        if not wanted or not any(t and (wanted in t or t in wanted) for t in titles):
            problems.append(f"{where}: column {col} is not titled '{value}' above the questions")
    return problems


def _check_target_column(
    snap: SheetSnapshot, col: str, question_idx: int, question_rows: list[int], filled_ok: bool = False
) -> list[str]:
    index = column_index_from_string(col)
    where = f"sheet '{snap.name}', column {col}"
    filled = 0
    for row in question_rows:
        question = snap.text(row, question_idx)
        cell = snap.text(row, index)
        if cell and question in cell:
            return [f"{where}: cell {col}{row} holds the question text"]
        area = snap.merged_at(row, index)
        if area is not None and area.contains(row, question_idx):
            return [f"{where}: cell {col}{row} is merged with the question cell"]
        filled += bool(cell)
    if not filled_ok and filled >= 2 and filled * 2 > len(question_rows):
        return [f"{where}: already holds data in {filled} of {len(question_rows)} question rows"]
    return []


def _check_allowed_values(
    snap: SheetSnapshot, short_col: str | None, allowed: list[str] | None, first: int, last: int
) -> tuple[list[str] | None, bool]:
    """The allowed values to keep and whether they agree with the sheet's drop-down list.

    Without a drop-down list over the column the values cannot be verified and are dropped. When the
    model gave none but the column has a list, the list itself is the allowed values.
    """
    if short_col is None:
        return None, True
    listed = snap.list_values(column_index_from_string(short_col), range(first, last + 1))
    if listed is None:
        return None, True
    if not allowed:
        return sorted(listed), True
    if {v.strip() for v in allowed} <= listed:
        return list(allowed), True
    return None, False
