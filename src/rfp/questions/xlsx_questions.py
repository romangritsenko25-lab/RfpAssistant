from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from rfp.models import Location, Question
from rfp.questions.layout_map import AnswerTargets


def extract_xlsx_questions(
    path: Path,
    sheet: str,
    question_col: str,
    answer_col: str,
    header_row: int = 1,
    number_col: str | None = None,
    *,
    last_row: int | None = None,
    skip_rows: frozenset[int] = frozenset(),
    targets: AnswerTargets | None = None,
) -> list[Question]:
    """Read one question per row below ``header_row``; rows with an empty question cell are skipped.

    Columns are letters (``"B"``). ``number_col`` optionally names a column holding the item number.
    ``skip_rows`` lists rows that are not questions (section titles); ``last_row`` ends the scan.
    ``targets`` adds the cells for the short answer, source and marks to every question's location.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet]
        question_idx = column_index_from_string(question_col)
        number_idx = column_index_from_string(number_col) if number_col else None
        questions: list[Question] = []
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, max_row=last_row), start=header_row + 1
        ):
            text = _cell_text(row, question_idx)
            if not text or row_number in skip_rows:
                continue
            number = _cell_text(row, number_idx) if number_idx else ""
            questions.append(
                Question(
                    qid=f"q{len(questions) + 1}",
                    text=text,
                    number=number or None,
                    location=Location(
                        sheet=sheet,
                        row=row_number,
                        answer_col=answer_col.upper(),
                        question_col=question_col.upper(),
                        short_col=targets.short_answer_col if targets else None,
                        allowed_values=targets.allowed_values if targets else None,
                        source_col=targets.source_col if targets else None,
                        choice_cols=targets.choice_cols if targets else None,
                    ),
                )
            )
        return questions
    finally:
        workbook.close()


def _cell_text(row: tuple, index: int) -> str:
    if index > len(row) or row[index - 1].value is None:
        return ""
    return str(row[index - 1].value).strip()
