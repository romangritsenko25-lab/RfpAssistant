"""Self-check of an exported file: reopen it and compare it with the source and with the answers.

A file that fails the check is never handed to the user. The check recomputes what the export
should have written (``exported_text``) and demands that the result differs from the source in
nothing else: same sheets, sizes, formulas, merged cells and drop-down lists for XLSX; same
paragraphs and table cells in the same order for DOCX, plus the answers where they belong.
"""

from collections.abc import Sequence
from pathlib import Path

from docx import Document as open_docx
from docx.table import Table
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from rfp.docx_blocks import iter_blocks
from rfp.export.formatting import exported_text
from rfp.export.targets import find_mark
from rfp.export.xlsx_export import AnswerHeader, extra_cells
from rfp.models import Answer, Question
from rfp.questions.docx_questions import unique_cells
from rfp.questions.placeholders import is_placeholder

MAX_PROBLEMS = 5


class ExportVerificationError(ValueError):
    """The exported file does not match the source and the answers; the message says where."""

    def __init__(self, problems: list[str]) -> None:
        self.problems = problems[:MAX_PROBLEMS]
        super().__init__(" | ".join(self.problems))


def verify_xlsx(
    source: Path,
    dest: Path,
    questions: list[Question],
    answers: dict[str, Answer],
    headers: Sequence[AnswerHeader] = (),
) -> None:
    before, after = load_workbook(source), load_workbook(dest)
    try:
        if before.sheetnames != after.sheetnames:
            raise ExportVerificationError(
                [f"sheets differ: expected {before.sheetnames}, got {after.sheetnames}"]
            )
        writes = _expected_writes(before, questions, answers, headers)
        problems: list[str] = []
        for name in before.sheetnames:
            problems += _check_sheet(before[name], after[name], writes)
        problems += _check_short_values(after, questions)
        if problems:
            raise ExportVerificationError(problems)
    finally:
        before.close()
        after.close()


def _check_short_values(after, questions: list[Question]) -> list[str]:
    """No short-answer cell may hold a value the sheet does not allow."""
    problems = []
    for question in questions:
        location = question.location
        if not (location.short_col and location.allowed_values and location.sheet and location.row):
            continue
        value = after[location.sheet][f"{location.short_col}{location.row}"].value
        if value is not None and str(value).strip() not in {v.strip() for v in location.allowed_values}:
            problems.append(
                f"sheet '{location.sheet}', cell {location.short_col}{location.row}: "
                f"'{value}' is not an allowed value"
            )
    return problems


def _expected_writes(workbook, questions, answers, headers) -> dict[tuple[str, int, int], tuple[str, str]]:
    """Address -> (expected value, what it is for)."""
    writes: dict[tuple[str, int, int], tuple[str, str]] = {}
    marks = {
        q.location.sheet: find_mark(workbook[q.location.sheet], q.location.choice_cols)
        for q in questions
        if q.location.sheet in workbook.sheetnames and q.location.choice_cols
    }
    for question in questions:
        answer = answers.get(question.qid)
        location = question.location
        if answer is None or location.sheet is None or location.row is None or location.answer_col is None:
            continue
        key = (location.sheet, location.row, column_index_from_string(location.answer_col))
        writes[key] = (exported_text(question, answer), f"answer for {question.qid}")
        for col, value in extra_cells(question, answer, marks.get(location.sheet, "")):
            extra = (location.sheet, location.row, column_index_from_string(col))
            writes[extra] = (value, f"short answer, source or mark for {question.qid}")
    for header in headers:
        key = (header.sheet, header.row, column_index_from_string(header.col))
        if workbook[header.sheet].cell(row=key[1], column=key[2]).value is None:
            writes[key] = (header.text, "the new answer column title")
    return writes


def _check_sheet(before, after, writes: dict[tuple[str, int, int], tuple[str, str]]) -> list[str]:
    name = before.title
    problems: list[str] = []
    own = {(r, c): item for (s, r, c), item in writes.items() if s == name}
    rows = max([before.max_row, *(r for r, _ in own)])
    cols = max([before.max_column, *(c for _, c in own)])
    if (after.max_row, after.max_column) != (rows, cols):
        problems.append(f"sheet '{name}': size is {after.max_row}x{after.max_column}, expected {rows}x{cols}")
    if _formulas(before) != _formulas(after):
        problems.append(f"sheet '{name}': the number of formulas changed")
    if sorted(str(r) for r in before.merged_cells.ranges) != sorted(
        str(r) for r in after.merged_cells.ranges
    ):
        problems.append(f"sheet '{name}': merged cells changed")
    if len(before.data_validations.dataValidation) != len(after.data_validations.dataValidation):
        problems.append(f"sheet '{name}': drop-down lists changed")
    old, new = _values(before), _values(after)
    for address in sorted(old.keys() | new.keys()):
        if address in own:
            continue
        if old.get(address) != new.get(address):
            problems.append(f"sheet '{name}', cell {_label(address)}: content changed")
    for address, (expected, purpose) in own.items():
        if (new.get(address) or "") != expected:
            problems.append(f"sheet '{name}': {purpose} is not at {_label(address)}")
    return problems


def _values(sheet) -> dict[tuple[int, int], object]:
    return {
        (cell.row, cell.column): cell.value
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }


def _formulas(sheet) -> int:
    return sum(1 for row in sheet.iter_rows() for cell in row if cell.data_type == "f")


def _label(address: tuple[int, int]) -> str:
    return f"{get_column_letter(address[1])}{address[0]}"


# --- DOCX ---------------------------------------------------------------------------------------


def verify_docx(source: Path, dest: Path, questions: list[Question], answers: dict[str, Answer]) -> None:
    before_blocks = list(iter_blocks(open_docx(str(source))))
    after_blocks = list(iter_blocks(open_docx(str(dest))))
    below: dict[int, list[str]] = {}
    in_cell: dict[tuple[int, int, int], list[str]] = {}
    for question in questions:
        answer = answers.get(question.qid)
        location = question.location
        if answer is None or location.block is None:
            continue
        text = exported_text(question, answer)
        if location.table_row is None:
            below.setdefault(location.block, []).append(text)
        else:
            row = location.answer_table_row if location.answer_table_row is not None else location.table_row
            col = location.answer_table_col if location.answer_table_col is not None else location.table_col
            in_cell.setdefault((location.block, row, col or 0), []).append(text)

    expected: list[str | None] = []  # paragraph text, or None for a table
    for index, block in enumerate(before_blocks):
        expected.append(None if isinstance(block, Table) else block.text)
        expected += list(reversed(below.get(index, [])))  # each insertion lands right under the block
    actual = [None if isinstance(block, Table) else block.text for block in after_blocks]
    problems: list[str] = []
    if len(expected) != len(actual):
        problems.append(f"the document has {len(actual)} blocks, expected {len(expected)}")
    else:
        for position, (want, got) in enumerate(zip(expected, actual, strict=True), 1):
            if want != got:
                problems.append(f"block {position}: text differs from the expected one")
    if problems:
        raise ExportVerificationError(problems)

    old_tables = [(i, b) for i, b in enumerate(before_blocks) if isinstance(b, Table)]
    new_tables = [b for b in after_blocks if isinstance(b, Table)]
    for (block_index, old), new in zip(old_tables, new_tables, strict=True):
        problems += _check_table(block_index, old, new, in_cell)
    if problems:
        raise ExportVerificationError(problems)


def _check_table(block_index: int, old: Table, new: Table, in_cell: dict) -> list[str]:
    if len(old.rows) != len(new.rows):
        return [f"table in block {block_index + 1}: the number of rows changed"]
    problems = []
    for row_index, (old_row, new_row) in enumerate(zip(old.rows, new.rows, strict=True)):
        old_cells, new_cells = unique_cells(old_row), unique_cells(new_row)
        if [c for c, _ in old_cells] != [c for c, _ in new_cells]:
            problems.append(f"table in block {block_index + 1}, row {row_index + 1}: cells changed")
            continue
        for (col, old_cell), (_, new_cell) in zip(old_cells, new_cells, strict=True):
            before = [p.text for p in old_cell.paragraphs]
            want = _expected_cell(before, in_cell.get((block_index, row_index, col), []))
            if want != [p.text for p in new_cell.paragraphs]:
                problems.append(
                    f"table in block {block_index + 1}, row {row_index + 1}, column {col + 1}: "
                    "text or answer differs from the expected one"
                )
    return problems


def _expected_cell(paragraphs: list[str], answers: list[str]) -> list[str]:
    if not answers:
        return paragraphs
    text = "\n".join(paragraphs)
    if text.strip() and is_placeholder(text):
        return list(answers)  # the template stub is replaced by the answer
    if not text.strip():
        return [answers[0], *paragraphs[1:], *answers[1:]]
    return [*paragraphs, *answers]
