from collections.abc import Sequence
from copy import copy
from pathlib import Path
from typing import NamedTuple

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import column_index_from_string

from rfp.export.formatting import exported_text
from rfp.export.targets import find_mark, mark_column, short_value, source_value
from rfp.models import Answer, Question

NO_DATA_FILL = PatternFill("solid", fgColor="FFF2CC")  # soft yellow: a person has to look here


class AnswerHeader(NamedTuple):
    """Title for an answer column the source file did not have; styled like the cell to its left."""

    sheet: str
    row: int
    col: str
    text: str


def export_xlsx(
    source: Path,
    questions: list[Question],
    answers: dict[str, Answer],
    dest: Path,
    answer_headers: Sequence[AnswerHeader] = (),
) -> None:
    """Write each answer into its answer cell of a copy of ``source``; nothing else is touched.

    Refuses (``ValueError``) to write into the question column or over a cell that already holds
    the question text, so a wrong column choice cannot destroy the questions. An answer of status
    ``no_data`` is written as a fixed note in the language of the question, on a yellow fill. When
    the sheet has them, the short value (only one the sheet allows), the source and the mark column
    of an answered question are written too.
    """
    workbook = load_workbook(source)
    marks = {
        q.location.sheet: find_mark(workbook[q.location.sheet], q.location.choice_cols)
        for q in questions
        if q.location.sheet in workbook.sheetnames and q.location.choice_cols
    }
    for question in questions:
        answer = answers.get(question.qid)
        if answer is None:
            continue
        location = question.location
        if location.sheet is None or location.row is None or location.answer_col is None:
            raise ValueError(f"Question {question.qid} has no XLSX location")
        if location.answer_col.upper() == (location.question_col or "").upper():
            raise ValueError(f"Question {question.qid}: the answer column is the question column")
        cell = workbook[location.sheet][f"{location.answer_col}{location.row}"]
        if isinstance(cell.value, str) and cell.value.strip() == question.text:
            raise ValueError(
                f"Cell {location.answer_col}{location.row} already holds the question text; "
                "it would be overwritten. Check the answer column."
            )
        cell.value = exported_text(question, answer)
        if cell.data_type == "f":  # an answer that starts with "=" must stay text, not become a formula
            cell.data_type = "s"
        if answer.status == "no_data":
            cell.fill = NO_DATA_FILL
        sheet = workbook[location.sheet]
        for col, value in extra_cells(question, answer, marks.get(location.sheet, "")):
            if col.upper() == (location.question_col or "").upper():
                raise ValueError(f"Question {question.qid}: column {col} is the question column")
            sheet[f"{col}{location.row}"].value = value
    for header in answer_headers:
        _write_header(workbook, header)
    workbook.save(dest)


def _write_header(workbook, header: AnswerHeader) -> None:
    sheet = workbook[header.sheet]
    cell = sheet[f"{header.col}{header.row}"]
    if cell.value is not None:
        return
    cell.value = header.text
    index = column_index_from_string(header.col)
    if index > 1:
        cell._style = copy(sheet.cell(row=header.row, column=index - 1)._style)


def extra_cells(question: Question, answer: Answer, mark: str) -> list[tuple[str, str]]:
    """(column, value) of the short answer, source and mark cells to write for this question."""
    location = question.location
    cells: list[tuple[str | None, str | None]] = [
        (location.short_col, short_value(location, answer)),
        (location.source_col, source_value(answer)),
        (mark_column(location, answer), mark),
    ]
    return [(col, value) for col, value in cells if col and value]
