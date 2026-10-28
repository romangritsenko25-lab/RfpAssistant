from pathlib import Path

from docx import Document as open_docx
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from rfp.docx_blocks import iter_blocks
from rfp.export.formatting import no_data_label, sources_line
from rfp.models import Answer, Question

ANSWER_STYLE = "RFP Answer"
_SOURCES_COLOR = RGBColor(0x59, 0x59, 0x59)
_SOURCES_SIZE = Pt(8)


def export_docx(source: Path, questions: list[Question], answers: dict[str, Answer], dest: Path) -> None:
    """Write each answer into a copy of ``source``: into the answer cell the question points to, or as a
    new paragraph right under the question. The question's own text is never changed."""
    doc = open_docx(str(source))
    _ensure_answer_style(doc)
    blocks = list(iter_blocks(doc))  # resolved up front: insertions shift later positions
    for question in questions:
        answer = answers.get(question.qid)
        if answer is not None:
            _write_answer(_answer_paragraph(doc, blocks, question), question, answer)
    doc.save(dest)


def export_docx_report(questions: list[Question], answers: dict[str, Answer], dest: Path) -> None:
    """Build a standalone DOCX of question, answer and sources; used for PDF inputs."""
    doc = open_docx()
    _ensure_answer_style(doc)
    doc.add_heading("Answers", level=1)
    for question in questions:
        answer = answers.get(question.qid)
        if answer is None:
            continue
        label = f"{question.number}. " if question.number else ""
        doc.add_paragraph().add_run(f"{label}{question.text}").bold = True
        _write_answer(doc.add_paragraph(), question, answer)
    doc.save(dest)


def _answer_paragraph(doc: DocxDocument, blocks: list, question: Question) -> Paragraph:
    location = question.location
    if location.block is None:
        raise ValueError(f"Question {question.qid} has no DOCX location")
    block = blocks[location.block]
    if isinstance(block, Table):
        cell = block.cell(location.table_row or 0, location.table_col or 0)
        if location.answer_table_row is not None and location.answer_table_col is not None:
            target = block.cell(location.answer_table_row, location.answer_table_col)
            if target._tc is cell._tc:
                raise ValueError(f"Question {question.qid}: the answer cell is the question cell")
            cell = target
        # An empty cell keeps its own (empty) paragraph; a cell with text gets a new one after it.
        return cell.add_paragraph() if cell.text.strip() else cell.paragraphs[0]
    new_p = OxmlElement("w:p")
    block._p.addnext(new_p)
    return Paragraph(new_p, doc)


def _write_answer(paragraph: Paragraph, question: Question, answer: Answer) -> None:
    paragraph.style = ANSWER_STYLE
    if answer.status == "no_data":
        paragraph.add_run(no_data_label(question)).italic = True
        return
    paragraph.add_run(answer.text)
    sources = sources_line(answer)
    if sources:
        run = paragraph.add_run()
        run.add_break()
        run.add_text(sources)
        run.font.size = _SOURCES_SIZE
        run.font.color.rgb = _SOURCES_COLOR


def _ensure_answer_style(doc: DocxDocument) -> None:
    if ANSWER_STYLE in doc.styles:
        return
    style = doc.styles.add_style(ANSWER_STYLE, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.paragraph_format.left_indent = Cm(0.8)
    style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
