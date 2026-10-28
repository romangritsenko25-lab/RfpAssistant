from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from rfp.documents.common import Parsed
from rfp.models import Section


def read_xlsx(path: Path) -> Parsed:
    """Turn every non-empty data row into a section ``Sheet / row N`` of ``header: value`` lines."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[Section] = []
    try:
        for sheet in workbook.worksheets:
            header: list[str] | None = None
            for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = ["" if value is None else str(value).strip() for value in row]
                if not any(cells):
                    continue
                if header is None:
                    header = cells
                    continue
                heading = f"{sheet.title} / row {number}"
                sections.append(Section(heading=heading, text=_row_text(header, cells), order=len(sections)))
        properties = workbook.properties
        stamp = properties.modified or properties.created
        return Parsed(
            sections=sections,
            title=properties.title or None,
            doc_id=properties.identifier or None,
            effective_date=stamp.date() if stamp else None,
        )
    finally:
        workbook.close()


def _row_text(header: list[str], cells: list[str]) -> str:
    lines = []
    for index, value in enumerate(cells):
        if value:
            name = header[index] if index < len(header) else ""
            lines.append(f"{name or get_column_letter(index + 1)}: {value}")
    return "\n".join(lines)
