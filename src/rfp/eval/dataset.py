"""Loading of an evaluation set from ``data/sets/<name>/``."""

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from rfp.eval.metrics import GoldItem
from rfp.models import Question
from rfp.questions import extract_docx_questions, extract_pdf_questions, extract_xlsx_questions

SETS_DIR = Path("data/sets")


def set_dir(name: str, root: Path = SETS_DIR) -> Path:
    return root / name


def load_gold(folder: Path) -> tuple[list[GoldItem], str]:
    """Read ``gold.jsonl``; return the items and the set version (first 12 hex chars of its SHA-256)."""
    raw = (folder / "gold.jsonl").read_bytes()
    items = [
        GoldItem.model_validate(json.loads(line)) for line in raw.decode("utf-8").splitlines() if line.strip()
    ]
    return items, hashlib.sha256(raw).hexdigest()[:12]


def extract_set_questions(folder: Path, files: set[str]) -> dict[str, list[Question]]:
    """Run the production question extraction on each source file of the set."""
    return {name: _extract(folder / name) for name in sorted(files)}


def _extract(path: Path) -> list[Question]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        sheet, question_col, answer_col = _locate_columns(path)
        return extract_xlsx_questions(path, sheet, question_col, answer_col)
    if suffix == ".docx":
        return extract_docx_questions(path)
    if suffix == ".pdf":
        return extract_pdf_questions(path)
    raise ValueError(f"unsupported question file: {path.name}")


def _locate_columns(path: Path) -> tuple[str, str, str]:
    """Find the first sheet and the letters of the ``Question`` and ``Answer`` header cells."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        headers = {
            str(cell.value).strip().lower(): cell.column_letter
            for cell in next(sheet.iter_rows(min_row=1, max_row=1))
            if cell.value
        }
        return sheet.title, headers["question"], headers["answer"]
    except KeyError as error:
        raise ValueError(f"{path.name}: header row needs 'Question' and 'Answer' columns") from error
    finally:
        workbook.close()
